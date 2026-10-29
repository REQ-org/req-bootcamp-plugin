"""Print the ground-truth placeholder catalog for the REQ IC memo template.

The `references/layouts.md` doc in this skill is a curated, human-written summary of
this same data, annotated with "when to use" guidance from a real IC memo. It can
drift out of sync if the template is ever updated. This script re-derives the raw
facts (layout names, placeholder idx/type/position) straight from the .pptx, so you
can always check the reference doc against ground truth, or look up a layout that
isn't discussed in prose there.

Usage:
    python scripts/dump_layouts.py                         # dumps the bundled template
    python scripts/dump_layouts.py /path/to/other.pptx      # dumps any other deck
    python scripts/dump_layouts.py --layout "Agenda"        # just one layout, by name (substring match)
    python scripts/dump_layouts.py --index 7                # just one layout, by 0-based index
"""

import argparse
import re
import sys
from pathlib import Path

from pptx import Presentation

DEFAULT_TEMPLATE = Path(__file__).parent.parent / "assets" / "REQ_Template.pptx"

PH_TAG_RE = re.compile(r"<p:ph\b[^>]*/>|<p:ph\b[^>]*>.*?</p:ph>", re.DOTALL)
IDX_ATTR_RE = re.compile(r'idx="(\d+)"')


def raw_ph_tags_by_idx(layout):
    """Map idx -> the literal <p:ph .../> tag text from this layout's XML.

    This is ground truth for the `type` attribute, which matters MORE than idx:
    a placeholder's type in your new slide must match the layout's exactly, or
    PowerPoint/LibreOffice silently drops position inheritance (the shape collapses
    to a fallback position, overlapping everything else) instead of erroring.
    validate.py will NOT catch this — it's a rendering-only defect. Copy the tag
    below verbatim into your slide's <p:nvPr> rather than reconstructing it from
    the python-pptx type name printed alongside it.
    """
    xml = layout.part.blob.decode("utf-8")
    out = {}
    for m in PH_TAG_RE.finditer(xml):
        tag = m.group(0)
        idx_m = IDX_ATTR_RE.search(tag)
        idx = int(idx_m.group(1)) if idx_m else 0  # idx defaults to 0 when absent (title)
        out[idx] = tag
    return out


def describe_layout(layout, index):
    print(f"--- Layout {index}: {layout.name}  (slideLayout{index + 1}.xml) ---")
    raw_tags = raw_ph_tags_by_idx(layout)
    for ph in layout.placeholders:
        pos = ""
        if ph.left is not None:
            pos = (
                f"pos=({ph.left/914400:.2f},{ph.top/914400:.2f},"
                f"{ph.width/914400:.2f},{ph.height/914400:.2f})in"
            )
        raw = raw_tags.get(ph.placeholder_format.idx, "")
        print(f"    idx={ph.placeholder_format.idx:<3} type={str(ph.placeholder_format.type):<20} "
              f"name='{ph.name}' {pos}")
        print(f"        raw tag (copy verbatim): {raw}")
    for shp in layout.shapes:
        if not shp.is_placeholder:
            print(f"    [non-placeholder, renders automatically — do not recreate] "
                  f"type={shp.shape_type} name='{shp.name}'")


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("pptx", nargs="?", default=str(DEFAULT_TEMPLATE), help="Path to a .pptx (defaults to the bundled REQ template)")
    parser.add_argument("--layout", help="Only show layouts whose name contains this substring (case-insensitive)")
    parser.add_argument("--index", type=int, help="Only show the layout at this 0-based index")
    args = parser.parse_args()

    prs = Presentation(args.pptx)
    if len(prs.slide_masters) != 1:
        print(f"Note: {len(prs.slide_masters)} slide masters found; using the first.", file=sys.stderr)
    master = prs.slide_masters[0]

    print(f"Slide size: {prs.slide_width/914400:.3f}in x {prs.slide_height/914400:.3f}in")
    print(f"{len(master.slide_layouts)} layouts available.\n")

    for i, layout in enumerate(master.slide_layouts):
        if args.index is not None and i != args.index:
            continue
        if args.layout and args.layout.lower() not in layout.name.lower():
            continue
        describe_layout(layout, i)


if __name__ == "__main__":
    main()
