from build_part1 import *
from openpyxl.chart import BarChart, LineChart, Reference, Series
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.text import Paragraph, ParagraphProperties, CharacterProperties
from openpyxl.drawing.spreadsheet_drawing import TwoCellAnchor, AnchorMarker
from openpyxl.formatting.rule import ColorScaleRule, CellIsRule
from openpyxl.styles import Alignment, Font, Border, Side

TEN = [(7, False), (8, False), (9, False), (10, False), (11, True)]   # rent-roll rows, is Oslo kommune
TNAMES = ["Rema 1000 – butikk", "Rema 1000 – garasje", "Telenor – basestasjon", "Telia – basestasjon", "Oslo kommune – 28 omsorgsboliger"]
def rrf(col, row): return f"'Rent roll'!${col}${row}"

# =====================================================================================
# CASH-FLOW BLOCK WRITER (used on CF and in the sensitivity engine)
# =====================================================================================
def write_block(ws, top, scen, cpi, title):
    B = scen == "B"
    R_ = {}
    section(ws, top, "B", "O", title)
    r = top + 1; subhead(ws, r, "Leieinntekter"); r += 1
    YS = lambda c: f"DATE({c}$6,1,1)"
    YE = lambda c: f"DATE({c}$6,12,31)"
    D = lambda c: f"({YE(c)}-{YS(c)}+1)"
    idxf = lambda c, f: f"(1+KPIstart*{f})*(1+{cpi}*{f})^({c}$5-1)"
    idxm = lambda c: f"(1+KPIstart)*(1+{cpi})^({c}$5-1)"
    soldfrac = lambda c: f"MAX(0,{YE(c)}-MAX({YS(c)},SalgDato)+1)/{D(c)}"
    tr = []
    for (row, ok), nm_ in zip(TEN, TNAMES):
        base, f, exp, mkt, vac = rrf("F", row), rrf("H", row), rrf("I", row), rrf("O", row), rrf("P", row)
        def fn(c, i, base=base, f=f, exp=exp, mkt=mkt, vac=vac, ok=ok):
            cfrac = f"MAX(0,MIN({YE(c)},{exp})-{YS(c)}+1)/{D(c)}"
            mfrac = f"MAX(0,{YE(c)}-MAX({YS(c)},EDATE({exp},{vac})+1)+1)/{D(c)}"
            if B and ok:
                return f"={cfrac}*{base}*{idxf(c, f)}"
            return f"={cfrac}*{base}*{idxf(c, f)}+{mfrac}*{mkt}*{idxm(c)}"
        tsrow(ws, r, "  " + nm_, fn); tr.append(r); r += 1
    R_["gross"] = r; tsrow(ws, r, "Brutto kontraktsleie", lambda c, i: f"=SUM({c}{tr[0]}:{c}{tr[-1]})", bold=True); r += 1
    R_["led"] = r; tsrow(ws, r, "Strukturell ledighet", lambda c, i: f"=-{c}{R_['gross']}*StruktLed*({c}$5>=StruktFra)"); r += 1
    R_["net"] = r; tsrow(ws, r, "Netto leieinntekter", lambda c, i: f"={c}{R_['gross']}+{c}{R_['led']}", bold=True); total_row(ws, r, "B", "O"); r += 1
    R_["eier"] = r; tsrow(ws, r, "Eierkost", lambda c, i: f"=-{c}{R_['net']}*Eierkost"); r += 1
    R_["es"] = r
    tsrow(ws, r, "Eiendomsskatt", (lambda c, i: f"=-EiendSkatt*{idxm(c)}*(1-BoligAndel*{soldfrac(c)})") if B else (lambda c, i: f"=-EiendSkatt*{idxm(c)}")); r += 1
    R_["noi"] = r; tsrow(ws, r, "NOI", lambda c, i: f"={c}{R_['net']}+{c}{R_['eier']}+{c}{R_['es']}", bold=True); total_row(ws, r, "B", "O", PALEGRN); r += 1
    subhead(ws, r, "Forvaltning"); r += 1
    R_["am"] = r
    tsrow(ws, r, "AM-kostnad" + (" (ikke i RREF)" if REFP else ""), (lambda c, i: f"=-AMpct*Kjopspris*(1-BoligAndel*{soldfrac(c)})") if B else (lambda c, i: "=-AMpct*Kjopspris")); r += 1
    R_["tf"] = r; tsrow(ws, r, "Teknisk forvaltning", lambda c, i: f"=-TFkr*{idxm(c)}"); r += 1
    R_["ff"] = r; tsrow(ws, r, "Forretningsførsel", lambda c, i: f"=-{c}{R_['net']}*FFpct"); r += 1
    R_["ebitda"] = r; tsrow(ws, r, "EBITDA", lambda c, i: f"={c}{R_['noi']}+{c}{R_['am']}+{c}{R_['tf']}+{c}{R_['ff']}", bold=True); total_row(ws, r, "B", "O"); r += 1
    subhead(ws, r, "Investeringer og kapitalhendelser"); r += 1
    def relet(c, i):
        terms = []
        for (row, ok) in TEN:
            if B and ok: continue
            exp, vac, cap, area = rrf("I", row), rrf("P", row), rrf("Q", row), rrf("E", row)
            terms.append(f"IF(YEAR(EDATE({exp},{vac}))={c}$6,-{cap}*{area}*{idxm(c)},0)")
        return "=" + "+".join(terms)
    R_["relet"] = r; tsrow(ws, r, "Capex ved re-letting", relet); r += 1
    R_["lek"] = r; tsrow(ws, r, "Capex-lekkasje" + (" (ikke i RREF)" if REFP else ""), lambda c, i: f"=-{c}{R_['net']}*CapexLek"); r += 1
    R_["opp"] = r; tsrow(ws, r, "Oppussing boligdel (seksjonssalg)", (lambda c, i: f"=IF(YEAR(OKUtlop)={c}$6,-OppussNom,0)") if B else 0); r += 1
    R_["seks"] = r; tsrow(ws, r, "Netto proveny seksjonssalg (etter skatt)", (lambda c, i: f"=IF(YEAR(SalgDato)={c}$6,SeksjonProveny,0)") if B else 0); r += 1
    subhead(ws, r, "Finansiering og skatt"); r += 1
    R_["ren"] = r; tsrow(ws, r, "Renter (bank + selgerkreditt, kontant)", lambda c, i: f"=Financing!{c}24", kind="link"); r += 1
    R_["avd"] = r; tsrow(ws, r, "Avdrag og innfrielser", lambda c, i: f"=Financing!{c}25", kind="link"); r += 1
    R_["str"] = r; tsrow(ws, r, "Skattemessig resultat", lambda c, i: f"={c}{R_['ebitda']}+Financing!{c}26-'Tax accounts'!{c}15"); r += 1
    pool = lambda c, i: "UnderskuddInn" if i == 0 else f"{YC[i-1]}{R_['str']+1}"
    R_["uf"] = r; tsrow(ws, r, "Fremførbart underskudd – UB", lambda c, i: f"=MAX(0,{pool(c, i)}-{c}{R_['str']})"); r += 1
    R_["skatt"] = r; tsrow(ws, r, "Løpende skatt", lambda c, i: f"=-SkattLop*MAX(0,{c}{R_['str']}-{pool(c, i)})"); r += 1
    R_["cfe"] = r
    tsrow(ws, r, "Kontantstrøm til egenkapital", lambda c, i: "=" + "+".join(f"{c}{R_[k]}" for k in ["ebitda", "relet", "lek", "opp", "seks", "ren", "avd", "skatt"]), bold=True)
    total_row(ws, r, "B", "O", PALEGRN); r += 1
    subhead(ws, r, "Memo – verdigrunnlag og nøkkeltall"); r += 1
    def rrrent(c, i, only_naering=False):
        terms = []
        for (row, ok) in TEN:
            if only_naering and ok: continue
            base, f, exp, mkt = rrf("F", row), rrf("H", row), rrf("I", row), rrf("O", row)
            after = "0" if (B and ok) else f"{mkt}*{idxm(c)}"
            terms.append(f"IF({YS(c)}<={exp},{base}*{idxf(c, f)},{after})")
        return "=" + "+".join(terms)
    R_["rr"] = r; tsrow(ws, r, "Annualisert leie per 1.1. – totalt", rrrent); r += 1
    R_["rrn"] = r; tsrow(ws, r, "Annualisert leie per 1.1. – næring", lambda c, i: rrrent(c, i, True)); r += 1
    adj = lambda c: f"(1-StruktLed*({c}$5>=StruktFra))*(1-Eierkost)"
    R_["rrnoi"] = r; tsrow(ws, r, "Annualisert NOI – totalt", lambda c, i: f"={c}{R_['rr']}*{adj(c)}+{c}{R_['es']}"); r += 1
    R_["rrnoin"] = r; tsrow(ws, r, "Annualisert NOI – næring", lambda c, i: f"={c}{R_['rrn']}*{adj(c)}+{c}{R_['es']}*IF({c}{R_['rr']}=0,0,{c}{R_['rrn']}/{c}{R_['rr']})"); r += 1
    R_["rrnoib"] = r; tsrow(ws, r, "Annualisert NOI – bolig", lambda c, i: f"={c}{R_['rrnoi']}-{c}{R_['rrnoin']}"); r += 1
    R_["ucf"] = r; tsrow(ws, r, "Kontantstrøm før finansiering (ubelånt)", lambda c, i: "=" + "+".join(f"{c}{R_[k]}" for k in ["ebitda", "relet", "lek", "opp", "seks"]) + f"-SkattLop*MAX(0,{c}{R_['ebitda']}-'Tax accounts'!{c}15)"); r += 1
    R_["coc"] = r; tsrow(ws, r, "Cash-on-cash (kontantstrøm / EK)", lambda c, i: f"={c}{R_['cfe']}/Egenkapital", F_PCT); r += 1
    R_["dpi"] = r; tsrow(ws, r, "DPI (akkumulert kontantstrøm / EK)", lambda c, i: f"={c}{R_['coc']}" if i == 0 else f"={YC[i-1]}{R_['dpi']}+{c}{R_['coc']}", F_PCT); r += 1
    line_rows(ws, top + 2, r - 1, "B", "O")
    for k in ["net", "noi", "ebitda", "cfe"]:
        total_row(ws, R_[k], "B", "O", PALEGRN if k in ("noi", "cfe") else PALE)
    R_["end"] = r
    return R_

# ---- CF sheet
cf = S["CF"]
ts_setup(cf, "CASH FLOW · 10 ÅRS KONTANTSTRØM (NOK)", "Scenario A og B er like frem til Oslo kommunes kontraktsutløp · kolonne 11 kun for annualisert NOI ved exit år 10")
A = write_block(cf, 8, "A", "KPI", "Scenario A · Oslo kommune re-leier ved utløp 2033")
Bk = write_block(cf, A["end"] + 2, "B", "KPI", "Scenario B · Oslo kommune fraflytter – boligdel seksjoneres og selges")
R["A"], R["B"] = A, Bk

# fill Overview formulas that needed CF rows
ov[f"C{R['ov_entry']}"] = f"=CF!E{A['noi']}/Kjopspris"
ov[f"C{R['ov_andel']}"] = f"=CF!E{A['rrnoib']}/CF!E{A['rrnoi']}"
ov[f"C{R['ov_eierkvm']}"] = f"=-CF!E{A['eier']}/ArealTot"

# ---- Tax accounts running tax rows
tsrow(tx, 19, "Skattemessig resultat – A", lambda c, i: f"=CF!{c}{A['str']}", kind="link")
tsrow(tx, 20, "Løpende skatt – A", lambda c, i: f"=CF!{c}{A['skatt']}", kind="link")
tsrow(tx, 21, "Skattemessig resultat – B", lambda c, i: f"=CF!{c}{Bk['str']}", kind="link")
tsrow(tx, 22, "Løpende skatt – B", lambda c, i: f"=CF!{c}{Bk['skatt']}", kind="link")
line_rows(tx, 19, 22, "B", "O")

# ---- Financing KPIs / covenants (scenario A)
from openpyxl.worksheet.formula import ArrayFormula
tsrow(fi, 30, "Verdi (annualisert NOI år+1 / exit yield)", lambda c, i: f"=INDEX(CF!$E${A['rrnoi']}:$O${A['rrnoi']},MIN(11,{c}$5+1))/ExitYield")
tsrow(fi, 31, "LTV bank", lambda c, i: f"={c}13/{c}30", F_PCT)
tsrow(fi, 32, "LTV total (bank + selgerkreditt)", lambda c, i: f"={c}27/{c}30", F_PCT)
tsrow(fi, 33, "ICR (EBITDA / rentekostnad)", lambda c, i: f'=IF({c}26=0,"–",CF!{c}{A["ebitda"]}/-{c}26)', F_X)
tsrow(fi, 34, "DSCR (EBITDA / renter + avdrag)", lambda c, i: f'=IF(({c}24+{c}12)=0,"–",CF!{c}{A["ebitda"]}/-({c}24+{c}12))', F_X)
tsrow(fi, 35, "Debt yield (NOI / gjeld IB)", lambda c, i: f'=IF(({c}11+{c}17)=0,"–",CF!{c}{A["noi"]}/({c}11+{c}17))', F_PCT)
tsrow(fi, 36, "Covenant LTV", lambda c, i: f'=IF({c}27=0,"–",IF({c}32<=CovLTV,"OK","BRUDD"))', None)
tsrow(fi, 37, "Covenant ICR", lambda c, i: f'=IF({c}26=0,"–",IF({c}33>=CovICR,"OK","BRUDD"))', None)
line_rows(fi, 30, 37, "B", "O")
for c in YC:
    for rr_ in (36, 37): fi[f"{c}{rr_}"].alignment = Alignment(horizontal="right")
fi.conditional_formatting.add("E36:O37", CellIsRule(operator="equal", formula=['"BRUDD"'], font=Font(color="C00000", bold=True), fill=fill("FDE2E2")))
put(fi, "B39", "Oppsummering (lang holdeperiode)", "label", bold=True)
kp = [("LTV bank ved kjøp", "=IF(Kjopspris=0,0,LaanT0/Kjopspris)", F_PCT, "LTVbank0"),
      ("LTV total ved kjøp", "=IF(Kjopspris=0,0,(LaanT0+SKBelop)/Kjopspris)", F_PCT, "LTVtot0"),
      ("All-in rente bank år 1", "=E10", F_PCT2, "AllIn1"),
      ("ICR år 1", "=E33", F_X, "ICR1"),
      ("Debt yield år 1", "=E35", F_PCT, "DY1"),
      ("Maks LTV total i holdeperioden", "=SUMPRODUCT(MAX(($E$5:$O$5<=ExitLang)*$E$32:$O$32))", F_PCT, "LTVmaks"),
      ("Gjeld ved exit (lang)", "=INDEX($E$27:$O$27,ExitLang)", F_NOK, "GjeldExit"),
      ("Antall covenant-brudd i holdeperioden", '=SUMPRODUCT(($E$5:$O$5<=ExitLang)*(($E$36:$O$36="BRUDD")+($E$37:$O$37="BRUDD")))', F_INT, "CovBrudd")]
for j, (lab, f, fmt, nm_) in enumerate(kp):
    put(fi, f"B{40+j}", lab, "label"); put(fi, f"E{40+j}", f, "calc", fmt); name(wb, nm_, "Financing", f"E{40+j}")
put(fi, f"B{40+len(kp)}", "Min ICR i holdeperioden", "label")
fi[f"E{40+len(kp)}"] = ArrayFormula(f"E{40+len(kp)}", '=IF(LaanT0+SKBelop=0,"–",MIN(IF(($E$5:$O$5<=ExitLang)*ISNUMBER($E$33:$O$33),$E$33:$O$33)))')
fi[f"E{40+len(kp)}"].number_format = F_X; fi[f"E{40+len(kp)}"].font = font(9); name(wb, "ICRmin", "Financing", f"E{40+len(kp)}")
line_rows(fi, 40, 40 + len(kp), "B", "E")

# =====================================================================================
# BOLIGVERDI
# =====================================================================================
bv = S["Boligverdi"]
setup(bv, {"A": 2, "B": 56, "C": 16, "D": 12, "E": 80}, zoom=90)
banner(bv, "BOLIGVERDI · RESIDUALVERDI, IMPLISITT PRIS OG SEKSJONSSALG", "Markedsstøtte for boligdelen dersom Oslo kommune ikke fornyer – grunnlag for scenario B", "B", "E")
section(bv, 5, "B", "E", "1 · Markedsdata – pris per kvm leilighet")
header_row(bv, 6, "B", "E", ["Kilde", "NOK/kvm", "Periode", "Kommentar"])
md = [("Komparable salg Aslakveien 16/18/33 (9 salg)", 92_461, "2023–2026", "Snitt i REQ arbeidsbok. Bygg-spesifikt – brukes som markedspris (MarkedKvm)."),
      ("Snittpris leilighet Røa, siste 12 mnd", 87_867, "mar. 2026", "Vårt Oslo / Eiendomsverdi"),
      ("Snittpris boliger Røa", 84_915, "sep. 2026", "Meglerprofil – indikasjon, alle boligtyper"),
      ("Vestre Aker, typisk intervall leiligheter", "80 000–100 000", "2025/26", "Boio bydelsoversikt")]
for j, (lab, v, per, note) in enumerate(md):
    r = 7 + j
    put(bv, f"B{r}", lab, "label"); put(bv, f"C{r}", v, "input", F_NOK, align="right")
    put(bv, f"D{r}", per, "note", align="center"); put(bv, f"E{r}", note, "note")
line_rows(bv, 7, 10, "B", "E"); name(wb, "MarkedKvm", "Boligverdi", "C7")

def bvrows(start, rows):
    r = start
    for lab, f, fmt, nm_, bold, note in rows:
        put(bv, f"B{r}", lab, "label", bold=bold); put(bv, f"C{r}", f, "calc" if str(f).startswith("=") else "input", fmt, bold=bold)
        if note: put(bv, f"E{r}", note, "note")
        if nm_: name(wb, nm_, "Boligverdi", f"C{r}")
        if bold: total_row(bv, r, "B", "E")
        r += 1
    line_rows(bv, start, r - 1, "B", "E")
    for rr_ in range(start, r):
        if bv[f"C{rr_}"].font.bold: total_row(bv, rr_, "B", "E")
    return r

section(bv, 12, "B", "E", "2 · Residualverdi boligdel i dag (2026-kr)")
r = bvrows(13, [
    ("Salgbart areal (BRA)", "=ArealSalg", F_NOK, None, False, ""),
    ("Egen pris/kvm i salgbar stand", "=EgenKvm", F_NOK, None, False, "Settes i Overview (foreløpig = markedspris)"),
    ("Rabatt egen pris vs. markedspris (utledet)", "=1-EgenKvm/MarkedKvm", F_PCT, "RabattEgen", False, "Utledes – settes ikke direkte"),
    ("Brutto verdi leiligheter", "=ArealSalg*EgenKvm", F_NOK, None, False, ""),
    ("+ P-plasser solgt med seksjonene", "=PPlasser*PVerdi", F_NOK, None, False, ""),
    ("− Oppussingscapex til salgbar stand", "=-CapexKvm*ArealSalg", F_NOK, None, False, ""),
    ("− Salgskostnad", "=-SalgKost*(C16+C17)", F_NOK, None, False, ""),
    ("Netto residualverdi (før skatt)", "=SUM(C16:C19)", F_NOK, "Residual0", True, "Markedsgulv for boligdelen"),
    ("Netto residualverdi per kvm", "=Residual0/ArealSalg", F_NOK, None, False, "")])

section(bv, 24, "B", "E", "3 · Implisitt pris for boligdelen i kjøpet")
r = bvrows(25, [
    ("Kjøpspris", "=Kjopspris", F_NOK, None, False, ""),
    ("− Verdi næringsdel (annualisert NOI næring år 1 / yield næring)", f"=-CF!E{A['rrnoin']}/YieldNaering", F_NOK, None, False, "Rema + telekom"),
    ("Implisitt verdi boligdel", "=C25+C26", F_NOK, "ImplBolig", True, ""),
    ("Implisitt pris per kvm bolig", "=ImplBolig/ArealSalg", F_NOK, "ImplKvm", False, ""),
    ("Rabatt vs. markedspris (komparable)", "=1-ImplKvm/MarkedKvm", F_PCT, "ImplRabatt", False, "Nedsidebeskyttelse ved konvertering"),
    ("Netto residualverdi i % av implisitt boligverdi", "=Residual0/ImplBolig", F_PCT, "Dekning", False, "> 100 % = residual dekker implisitt pris etter capex og salgskost"),
    ("Verdistøtte bottom-up (næring + bolig going concern)", f"=-C26+CF!E{A['rrnoib']}/CapBolig", F_NOK, "VerdiBottomUp", True, "Bolig kapitalisert til yield boligdel (going concern)"),
    ("Kjøpspris vs. bottom-up", "=Kjopspris/VerdiBottomUp-1", F_PCT, None, False, "Negativ = kjøp under bottom-up verdi")])
name(wb, "VerdiNaering", "Boligverdi", "C26")

section(bv, 35, "B", "E", "4 · Seksjonssalg – scenario B")
r = bvrows(36, [
    ("Kontraktsutløp Oslo kommune", "=OKUtlop", F_DATE, None, False, ""),
    ("Salgsdato seksjoner", "=SalgDato", F_DATE, None, False, ""),
    ("Salgsår", "=YEAR(SalgDato)", F_INT, "SalgAar", False, ""),
    ("Brutto salgsverdi (nominell)", "=(EgenKvm*ArealSalg+PPlasser*PVerdi)*(1+BoligVekst)^(SalgAar-(YEAR(Start)-1))", F_NOK, "BruttoSalg", False, "Prisvekst fra 2026-nivå"),
    ("− Salgskostnad", "=-SalgKost*BruttoSalg", F_NOK, None, False, ""),
    ("Indeks til utløpsår (capex)", "=(1+KPIstart)*(1+KPI)^(YEAR(OKUtlop)-YEAR(Start))", '0.000', "IdxUtlop", False, ""),
    ("Oppussingscapex (nominell, betales i utløpsåret)", "=CapexKvm*ArealSalg*IdxUtlop", F_NOK, "OppussNom", False, "Belastes CF scenario B i utløpsåret; aktiveres skattemessig"),
    ("Skattemessig saldo boligdel ved salg (allokert)", "=INDEX('Tax accounts'!$E$16:$O$16,MIN(11,MAX(1,SalgAar-YEAR(Start))))*BoligAndel", F_NOK, "SaldoSalg", False, "UB året før salg × boligandel"),
    ("Skattemessig verdi tomt (allokert)", "=TomtSkatt*BoligAndel", F_NOK, None, False, "Faktisk inngangsverdi fra saldoskjema"),
    ("Skattemessig gevinst", "=BruttoSalg+C40-C43-C44-OppussNom", F_NOK, None, False, "Salgssum − salgskost − saldo − tomt − aktivert oppussing"),
    ("− Gevinstskatt", "=-MAX(0,C45)*SkattSeksjon", F_NOK, None, False, "Full skatt i salgsåret (konservativt; gevinst- og tapskonto kan gi utsettelse)"),
    ("Netto proveny (etter salgskost og skatt)", "=BruttoSalg+C40+C46", F_NOK, "SeksjonProveny", True, "Inngår i CF scenario B i salgsåret")])

# =====================================================================================
# RETURNS
# =====================================================================================
rt = S["Returns"]
setup(rt, {"A": 2, "B": 60, **{get_column_letter(3 + i): 13.5 for i in range(11)}, "N": 2}, zoom=90)
banner(rt, "RETURNS · IRR, MoE, CASH-ON-CASH OG DPI", ("Belånt · brutto IRR på egenkapital før fondets honorar og carry" if REFP else "Ubelånt · IRR på egenkapital") + " (XIRR: kjøp 01.01, kontantstrømmer og exit 31.12) · EK etter skatteulempe og transaksjonskostnader", "B", "M")
section(rt, 5, "B", "H", "Exit-analyse per scenario og holdeperiode")
cases = [("A", "ExitKort"), ("A", "ExitMellom"), ("A", "ExitLang"), ("B", "ExitKort"), ("B", "ExitMellom"), ("B", "ExitLang")]
CC = ["C", "D", "E", "F", "G", "H"]
header_row(rt, 6, "B", "H", ["", "A · kort", "A · mellom", "A · lang", "B · kort", "B · mellom", "B · lang"])
def rtrow(r, label, fn, fmt=F_NOK, bold=False, kind="calc"):
    put(rt, f"B{r}", label, "label", bold=bold)
    for (sc, yn), c in zip(cases, CC):
        put(rt, f"{c}{r}", fn(sc, yn, c), kind, fmt, bold=bold)
rtrow(7, "Scenario", lambda s, y, c: s, None); [setattr(rt[f"{c}7"], "alignment", Alignment(horizontal="right")) for c in CC]
rtrow(8, "Holdeperiode (år)", lambda s, y, c: f"={y}", F_INT, kind="link")
rtrow(9, "Exit-dato", lambda s, y, c: f"=DATE(YEAR(Start)+{c}8-1,12,31)", F_DATE)
rtrow(10, "Boligdel ved exit (B: utleid / ledig / solgt)", lambda s, y, c: f'=IF({c}7="A","–",IF(SalgDato<={c}9,"Solgt",IF({c}9>=OKUtlop,"Ledig","Utleid")))', None)
[setattr(rt[f"{c}10"], "alignment", Alignment(horizontal="right")) for c in CC]
naer = lambda c: f'OR({c}10="Solgt",{c}10="Ledig")'
rtrow(11, "Annualisert NOI år N+1 (totalt · B ledig/solgt: næring)", lambda s, y, c: f"=INDEX(CF!$E${A['rrnoi']}:$O${A['rrnoi']},{c}8+1)" if s == "A" else
      f"=IF({naer(c)},INDEX(CF!$E${Bk['rrnoin']}:$O${Bk['rrnoin']},{c}8+1),INDEX(CF!$E${A['rrnoi']}:$O${A['rrnoi']},{c}8+1))")
rtrow(12, "Exit yield (inngangsyield · næring)", lambda s, y, c: "=ExitYield" if s == "A" else f"=IF({naer(c)},YieldNaering,ExitYield)", F_PCT2)
rtrow(13, "Verdi eiendom i drift", lambda s, y, c: f"={c}11/{c}12")
resid = lambda c: (f"((EgenKvm*ArealSalg+PPlasser*PVerdi)*(1+BoligVekst)^(YEAR({c}9)-(YEAR(Start)-1))*(1-SalgKost)"
                   f"-CapexKvm*ArealSalg*(1+KPIstart)*(1+KPI)^(YEAR({c}9)-YEAR(Start)))")
rat = lambda c: (f"({resid(c)}-SkattSeksjon*MAX(0,{resid(c)}-INDEX('Tax accounts'!$E$16:$O$16,{c}8)*BoligAndel-TomtSkatt*BoligAndel))")
gc = lambda c: f"INDEX(CF!$E${A['rrnoib']}:$O${A['rrnoib']},{c}8+1)/ExitYield"
rtrow(14, "Boligdel: residual etter skatt (ledig) / nedside-justering (utleid)", lambda s, y, c: "=0" if s == "A" else
      f'=IF({c}10="Ledig",{rat(c)},IF({c}10="Utleid",-MAX(0,{gc(c)}-{rat(c)}),0))')
rtrow(15, "Brutto exit-verdi", lambda s, y, c: f"={c}13+{c}14", bold=True)
rtrow(16, "Andel av skattemessig saldo", lambda s, y, c: f'=IF({c}10="Solgt",1-BoligAndel,1)', F_PCT)
rtrow(17, "Skattemessig saldo (UB exit-år × andel)", lambda s, y, c: f"=INDEX('Tax accounts'!$E$16:$O$16,{c}8)*{c}16")
rtrow(18, "Tomteverdi i skatteulempe (sats × exit-verdi)", lambda s, y, c: f"=TomtPct*{c}15")
rtrow(19, "− Skatteulempe", lambda s, y, c: f"=-SUsats*MAX(0,{c}15-{c}17-{c}18)")
rtrow(20, "− Divestment-kost", lambda s, y, c: f"=-Divest*{c}15")
rtrow(21, "− Restgjeld (bank + selgerkreditt)", lambda s, y, c: f"=-INDEX(Financing!$E$27:$O$27,{c}8)")
rtrow(22, "Netto exit-proveny til EK", lambda s, y, c: f"={c}15+{c}19+{c}20+{c}21", bold=True)
rtrow(24, "Egenkapital t0", lambda s, y, c: "=Egenkapital", kind="link")
cfe_sum = lambda s, c: f"=SUMPRODUCT((CF!$E$5:$O$5<={c}8)*CF!$E${(A if s=='A' else Bk)['cfe']}:$O${(A if s=='A' else Bk)['cfe']})"
rtrow(25, "Løpende kontantstrøm til EK i holdeperioden", lambda s, y, c: cfe_sum(s, c))
rtrow(26, "Total utbetalt til EK", lambda s, y, c: f"={c}25+{c}22", bold=True)
rtrow(27, "Gevinst", lambda s, y, c: f"={c}26-{c}24")
rtrow(28, "IRR", lambda s, y, c: f"=XIRR(C{35+CC.index(c)}:M{35+CC.index(c)},$C$34:$M$34)", F_PCT, bold=True)
rtrow(29, "MoE (equity multiple)", lambda s, y, c: f"={c}26/{c}24", F_X, bold=True)
rtrow(30, "Snitt cash-on-cash p.a.", lambda s, y, c: f"={c}25/{c}8/{c}24", F_PCT)
line_rows(rt, 7, 30, "B", "H")
for rr_ in (15, 22, 26): total_row(rt, rr_, "B", "H")
for rr_ in (28, 29): total_row(rt, rr_, "B", "H", PALEGRN)

section(rt, 32, "B", "M", "Kontantstrømmer til egenkapital (grunnlag for XIRR)")
header_row(rt, 33, "B", "M", ["Periode", "t0"] + [str(k) for k in range(1, 11)])
for k in range(1, 11):
    rt[f"{get_column_letter(3+k)}33"].value = k
put(rt, "B34", "Dato", "label", bold=True); put(rt, "C34", "=Start", "link", F_DATE)
for k in range(1, 11):
    put(rt, f"{get_column_letter(3+k)}34", f"=DATE(YEAR(Start)+{k}-1,12,31)", "calc", F_DATE)
labels = ["A · kort", "A · mellom", "A · lang", "B · kort", "B · mellom", "B · lang"]
for n, ((sc, yn), c) in enumerate(zip(cases, CC)):
    r = 35 + n
    put(rt, f"B{r}", f'="{labels[n]} ("&{c}8&" år)"', "label", align="left")
    put(rt, f"C{r}", "=-Egenkapital", "calc", F_NOK)
    cfer = (A if sc == "A" else Bk)["cfe"]
    for k in range(1, 11):
        col = get_column_letter(3 + k)
        put(rt, f"{col}{r}", f"=IF({col}$33<={c}$8,INDEX(CF!$E${cfer}:$O${cfer},{col}$33)+IF({col}$33={c}$8,{c}$22,0),0)", "calc", F_NOK)
line_rows(rt, 34, 40, "B", "M")

section(rt, 42, "B", "H", "Cash-on-cash og DPI (base case-definisjon som i pipeline deck)")
header_row(rt, 43, "B", "H", ["", "Kort", "Mellom", "Lang", "", "", ""])
for c, yn in zip(["C", "D", "E"], ["ExitKort", "ExitMellom", "ExitLang"]):
    rt[f"{c}43"].value = f'={yn}&" år"'
for j, (lab, blk, kind) in enumerate([("Cash-on-cash A (kontantstrøm år N / EK)", A, "coc"), ("DPI A (akkumulert / EK)", A, "dpi"),
                                       ("Cash-on-cash B", Bk, "coc"), ("DPI B", Bk, "dpi")]):
    r = 44 + j
    put(rt, f"B{r}", lab, "label")
    for c, yn in zip(["C", "D", "E"], ["ExitKort", "ExitMellom", "ExitLang"]):
        put(rt, f"{c}{r}", f"=INDEX(CF!$E${blk[kind]}:$O${blk[kind]},{yn})", "calc", F_PCT)
line_rows(rt, 44, 47, "B", "H")

section(rt, 49, "B", "H", "Diagramgrunnlag")
header_row(rt, 50, "B", "F", ["Holdeperiode", "IRR · A", "IRR · B", "MoE · A", "MoE · B"])
for j, (ca, cb) in enumerate([("C", "F"), ("D", "G"), ("E", "H")]):
    r = 51 + j
    put(rt, f"B{r}", f'={ca}8&" år"', "calc", align="left")
    put(rt, f"C{r}", f"={ca}28", "calc", F_PCT); put(rt, f"D{r}", f"={cb}28", "calc", F_PCT)
    put(rt, f"E{r}", f"={ca}29", "calc", F_X); put(rt, f"F{r}", f"={cb}29", "calc", F_X)
line_rows(rt, 51, 53, "B", "F")

section(rt, 55, "B", "M", "Ubelånt avkastning (eiendomsnivå, før finansiering)")
header_row(rt, 56, "B", "M", ["Periode", "t0"] + [str(k) for k in range(1, 11)])
for k in range(1, 11): rt[f"{get_column_letter(3+k)}56"].value = k
put(rt, "B57", "Dato", "label", bold=True); put(rt, "C57", "=Start", "link", F_DATE)
for k in range(1, 11): put(rt, f"{get_column_letter(3+k)}57", f"=DATE(YEAR(Start)+{k}-1,12,31)", "calc", F_DATE)
for n, ((sc, yn), c) in enumerate(zip(cases, CC)):
    r = 58 + n
    put(rt, f"B{r}", f'="{labels[n]} ("&{c}8&" år)"', "label", align="left")
    put(rt, f"C{r}", "=-(Kjopesum+DDSPA+WC+Megler+KjopsHon)", "calc", F_NOK)
    ur = (A if sc == "A" else Bk)["ucf"]
    for k in range(1, 11):
        col = get_column_letter(3 + k)
        put(rt, f"{col}{r}", f"=IF({col}$56<={c}$8,INDEX(CF!$E${ur}:$O${ur},{col}$56)+IF({col}$56={c}$8,{c}$15+{c}$19+{c}$20,0),0)", "calc", F_NOK)
put(rt, "B64", "IRR ubelånt", "label", bold=True); put(rt, "B65", "MoE ubelånt", "label", bold=True)
for n, c in enumerate(CC):
    put(rt, f"{c}64", f"=XIRR(C{58+n}:M{58+n},$C$57:$M$57)", "calc", F_PCT, bold=True)
    put(rt, f"{c}65", f"=SUM(D{58+n}:M{58+n})/-C{58+n}", "calc", F_X, bold=True)
line_rows(rt, 57, 63, "B", "M"); total_row(rt, 64, "B", "H", PALEGRN); total_row(rt, 65, "B", "H", PALEGRN)

# ---- Avkastningsbro (base case: A · lang, kolonne E i exit-analysen)
section(rt, 67, "B", "H", "Avkastningsbro · base case (A · lang holdeperiode)")
header_row(rt, 68, "B", "J", ["Komponent", "NOK", "Akk. start", "Akk. slutt", "Diagram: base", "Diagram: økning", "Diagram: reduksjon", "Diagram: total", "Diagram: etikett"], height=28)
SHORT = ["Egenkapital", "Drift", "Renter", "Skatt", "Verdiendring", "Skatteulempe / just.", "Transaksjonskost", "Totalt til EK"]
hold = lambda row: f"SUMPRODUCT((CF!$E$5:$O$5<=ExitLang)*CF!$E${row}:$O${row})"
br = [("Egenkapital investert", "=Egenkapital", "T"),
      ("Driftskontantstrøm (NOI − forvaltning − capex)", "=" + "+".join(hold(A[k]) for k in ["ebitda", "relet", "lek", "opp", "seks"]), "D"),
      ("Renter (inkl. kapitalisert)", "=SUMPRODUCT((Financing!$E$5:$O$5<=ExitLang)*Financing!$E$26:$O$26)", "D"),
      ("Løpende skatt", "=" + hold(A["skatt"]), "D"),
      ("Verdiendring eiendom (exit-verdi − eiendomsverdi)", "=E15-Kjopspris", "D"),
      ("Skatteulempe netto og kjøpesumsjusteringer", "=SkatteulempeInn+E19-JustKjopesum", "D"),
      ("Transaksjonskostnader (kjøp og salg)", "=-(DDSPA+WC+Megler+KjopsHon+EtabKost*LaanT0)+E20", "D"),
      ("Totalt utbetalt til egenkapital", None, "T")]
for j, (lab, f, kind) in enumerate(br):
    r = 69 + j
    put(rt, f"B{r}", lab, "label", bold=(kind == "T")); put(rt, f"J{r}", SHORT[j], "text", align="right")
    if j == len(br) - 1:
        f = f"=SUM(C69:C{r-1})"
    put(rt, f"C{r}", f, "calc", F_NOK, bold=(kind == "T"))
    if kind == "T":
        put(rt, f"D{r}", 0, "calc", F_NOK); put(rt, f"E{r}", f"=C{r}", "calc", F_NOK)
        put(rt, f"F{r}", 0, "calc", F_NOK); put(rt, f"G{r}", 0, "calc", F_NOK); put(rt, f"H{r}", 0, "calc", F_NOK)
        put(rt, f"I{r}", f"=C{r}", "calc", F_NOK)
    else:
        put(rt, f"D{r}", f"=E{r-1}", "calc", F_NOK); put(rt, f"E{r}", f"=D{r}+C{r}", "calc", F_NOK)
        put(rt, f"F{r}", f"=MAX(0,MIN(D{r},E{r}))", "calc", F_NOK)
        put(rt, f"G{r}", f"=MAX(0,C{r})", "calc", F_NOK); put(rt, f"H{r}", f"=MAX(0,-C{r})", "calc", F_NOK)
        put(rt, f"I{r}", 0, "calc", F_NOK)
for r in (69, 69 + len(br) - 1): total_row(rt, r, "B", "J", PALEGRN)
line_rows(rt, 70, 69 + len(br) - 2, "B", "J")
put(rt, "B78", "Kontroll: bro = total utbetalt til EK (Returns E26)", "note"); put(rt, "C78", "=C76-E26", "calc", '#,##0;(#,##0);"–"')
put(rt, "B79", "Avdrag er nøytrale (betales løpende, gir tilsvarende lavere restgjeld ved exit) og vises derfor ikke.", "note")
for r in range(69, 69 + len(br)):
    rt[f"F{r}"].number_format = '#,##0.0,,"m";;;'
    rt[f"G{r}"].number_format = '+#,##0.0,,"m";;;'
    rt[f"H{r}"].number_format = '-#,##0.0,,"m";;;'
    rt[f"I{r}"].number_format = '#,##0.0,,"m";;;'
R["rt_bridge"] = (69, 69 + len(br) - 1)

rt.freeze_panes = "C7"

import pickle
pickle.dump(R, open("./R.pkl", "wb"))
