from build_part3 import *
from build_part3 import _txt

sd = S["Sammendrag"]
COLS = [get_column_letter(ci("B") + k) for k in range(18)]   # B..S
setup(sd, {"A": 1.5, **{c: 9.6 for c in COLS}, "T": 1.5}, zoom=90)

def fillrange(c1, r1, c2, r2, color):
    for rr_ in range(r1, r2 + 1):
        for cc in range(ci(c1), ci(c2) + 1):
            sd.cell(row=rr_, column=cc).fill = fill(color)

def merge(c1, r1, c2, r2=None):
    r2 = r2 or r1
    if (c1, r1) != (c2, r2): sd.merge_cells(f"{c1}{r1}:{c2}{r2}")
    return sd[f"{c1}{r1}"]

def txt(c1, r1, c2, value, size=9, bold=False, color=TXT, align="left", fmt=None, r2=None, italic=False, wrap=False, valign="center"):
    c = merge(c1, r1, c2, r2); c.value = value
    c.font = Font(name=HEAD if bold else BODY, size=size, bold=bold, color=color, italic=italic)
    c.alignment = Alignment(horizontal=align, vertical=valign, wrap_text=wrap, indent=1 if align == "left" else 0)
    if fmt: c.number_format = fmt
    return c

def sec(c1, c2, r, text):
    section(sd, r, c1, c2, text, size=9.5)

def tile(r, c1, c2, value, fmt, label, size=17):
    fillrange(c1, r, c2, r + 1, PALE)
    txt(c1, r, c2, value, size=size, bold=True, color=NAVY, fmt=fmt, valign="bottom")
    txt(c1, r + 1, c2, label.upper(), size=7, bold=True, color=GREYTXT, valign="top")
    for rr_ in (r, r + 1):
        sd[f"{c1}{rr_}"].border = Border(left=Side(style="thick", color=WHITE))
        sd[f"{c2}{rr_}"].border = Border(right=Side(style="thick", color=WHITE))

# ---- banner
sd.row_dimensions[1].height = 8
sd.row_dimensions[2].height = 34; sd.row_dimensions[3].height = 18
fillrange("B", 2, "S", 2, NAVY); fillrange("B", 3, "S", 3, DEEP)
txt("B", 2, "N", "ASLAKVEIEN 18", size=18, bold=True, color=WHITE)
txt("O", 2, "S", "KONFIDENSIELT", size=8, bold=True, color=WHITE, align="right")
txt("B", 3, "N", "Oslo (Røa) · Eldre-/omsorgsbolig og nærsenter · " + ("REQ Real Estate Fund II · belånt analyse" if REFP else "REQ Evergreen Income Fund"), size=9.5, color=WHITE)
txt("O", 3, "S", "Underwriting-modell · sep. 2026", size=8, color=WHITE, align="right")
sd.row_dimensions[4].height = 8

# ---- KPI tiles
for rr_, h in [(5, 30), (6, 16), (7, 5), (8, 30), (9, 16)]:
    sd.row_dimensions[rr_].height = h
grossA = f"CF!E{A['gross']}"
T1 = [("Kjopspris", '#,##0,,"m"', "Kjøpspris (NOK)"), ("EntryYield", '0.0%', "Netto yield på kjøp"),
      ("WAULT", '0.0" år"', "WAULT (uten opsjon)"), ("Ledighet", '0.0%', "Ledighet i bygget"),
      ("Kjopspris/ArealTot", '#,##0', "Pris/kvm BTA (NOK)"), (grossA, '#,##0.0,,"m"', "Brutto leie år 1 (NOK)")]
T2 = [("ArealTot", '#,##0" kvm"', "Areal BTA"), (f"{grossA}/ArealTot", '#,##0', "Leie/kvm år 1 (NOK)"),
      (f"CF!E{A['noi']}", '#,##0.0,,"m"', "NOI år 1 (NOK)"), ("Egenkapital", '#,##0.0,,"m"', "Egenkapital t0 (NOK)"),
      ("OffAndel", '0.0%', "Offentlig andel av leie"), ("WAULTopt", '0.0" år"', "WAULT inkl. opsjoner")]
if REFP:
    T2 = [("ArealTot", '#,##0" kvm"', "Areal BTA"), (f"CF!E{A['noi']}", '#,##0.0,,"m"', "NOI år 1 (NOK)"),
          ("Egenkapital", '#,##0.0,,"m"', "Egenkapital t0 (NOK)"), ("LTVtot0", '0.0%', "LTV ved kjøp (total)"),
          ("AllIn1", '0.00%', "All-in rente år 1"), ("ICR1", '0.00"x"', "ICR år 1")]
spans = [("B", "D"), ("E", "G"), ("H", "J"), ("K", "M"), ("N", "P"), ("Q", "S")]
for (f, fmt, lab), (a, b) in zip(T1, spans): tile(5, a, b, f"={f}", fmt, lab)
for (f, fmt, lab), (a, b) in zip(T2, spans): tile(8, a, b, f"={f}", fmt, lab)
sd.row_dimensions[10].height = 10

# ---- base case (left) and returns table (right)
sec("B", "J", 11, f"Base case · {LANG} års hold · scenario A" + (" · belånt brutto" if REFP else ""))
sec("L", "S", 11, "Avkastning per holdeperiode og scenario")
for rr_, h in [(12, 28), (13, 14), (14, 16), (15, 24), (16, 13), (17, 16), (18, 24), (19, 13)]:
    sd.row_dimensions[rr_].height = h
tile(12, "B", "D", "=Returns!E28", '0.0%', "IRR på EK", size=20)
tile(12, "E", "G", "=Returns!E29", '0.00"x"', "Equity multiple", size=20)
tile(12, "H", "J", "=Returns!E27", '#,##0.0,,"m"', "Gevinst (NOK)", size=20)
txt("B", 14, "J", "CASH-ON-CASH (KONTANTSTRØM ÅR N / EK)", size=7.5, bold=True, color=DEEP, valign="bottom")
txt("B", 17, "J", "NETTO INVESTOR ETTER HONORAR OG CARRY (SHA 10.2)" if REFP else "DPI (AKKUMULERT UTBYTTE / INNSKUTT EK)", size=7.5, bold=True, color=DEEP, valign="bottom")
for (a, b), c, yn in zip([("B", "D"), ("E", "G"), ("H", "J")], ["C", "D", "E"], ["ExitKort", "ExitMellom", "ExitLang"]):
    tile(15, a, b, f"=Returns!{c}44", '0.0%', "år", size=13); sd[f"{a}16"].value = f'={yn}&" ÅR"'
    if not REFP:
        tile(18, a, b, f"=Returns!{c}45", '0.0%', "år", size=13); sd[f"{a}19"].value = f'={yn}&" ÅR"'
if REFP:
    tile(18, "B", "D", "=Utbyttefordeling!F9", '0.0%', "Netto IRR investor", size=13)
    tile(18, "E", "G", "=Utbyttefordeling!G9", '0.00"x"', "Netto MoE investor", size=13)
    tile(18, "H", "J", "=Utbyttefordeling!H9", '#,##0.0,,"m"', "Carried interest (NOK)", size=13)

# right table
for c in "LMNOPQRS":
    for rr_ in (12, 13):
        sd[f"{c}{rr_}"].fill = fill(DEEP)
hdr = lambda c1, r1, c2, v, r2=None: txt(c1, r1, c2, v, size=8, bold=True, color=WHITE, align="center", r2=r2, wrap=True)
hdr("L", 12, "L", "Hold", r2=13)
hdr("M", 12, "P", "A · Oslo kommune re-leier")
hdr("Q", 12, "S", "B · Fraflytting → seksjonssalg")
for c1, c2, v in [("M", "N", "Exit-verdi"), ("O", "O", "IRR"), ("P", "P", "MoE"), ("Q", "Q", "IRR"), ("R", "R", "MoE"), ("S", "S", "Exit-verdi")]:
    hdr(c1, 13, c2, v)
sd["P12"].border = Border(right=Side(style="thin", color=WHITE))
for n, (ca, cb) in enumerate([("C", "F"), ("D", "G"), ("E", "H")]):
    rr_ = 14 + n
    sd.row_dimensions[rr_].height = 20
    txt("L", rr_, "L", f'={ca}8&" år"'.replace(f"={ca}8", f"=Returns!{ca}8"), size=9, bold=True, align="center")
    txt("M", rr_, "N", f"=Returns!{ca}15", fmt='#,##0.0,,"m"', align="center")
    txt("O", rr_, "O", f"=Returns!{ca}28", fmt='0.0%', bold=True, color=NAVY, align="center")
    txt("P", rr_, "P", f"=Returns!{ca}29", fmt='0.00"x"', align="center")
    txt("Q", rr_, "Q", f"=Returns!{cb}28", fmt='0.0%', bold=True, color=NAVY, align="center")
    txt("R", rr_, "R", f"=Returns!{cb}29", fmt='0.00"x"', align="center")
    txt("S", rr_, "S", f"=Returns!{cb}15", fmt='#,##0.0,,"m"', align="center")
    for c in "LMNOPQRS":
        sd[f"{c}{rr_}"].border = Border(bottom=THIN)
        if n == 2: sd[f"{c}{rr_}"].fill = fill(PALEGRN)
txt("L", 17, "S", (("Belånt brutto IRR på EK (før fondets honorar og carry), etter 22 % skatt, skatteulempe og transaksjonskost (XIRR). " if REFP else "IRR på egenkapital etter skatteulempe og transaksjonskostnader (XIRR). ") + "Exit-verdi = annualisert NOI år N+1 / exit yield. "
                   "Scenario B, exit før utløp: kjøper priser boligdelen til laveste av leieverdi og residual etter gevinstskatt (nedside). "
                   "Etter utløp: seksjoner selges med gevinstskatt; næringsdel selges som aksjeselskap."), size=7.5, color=GREYTXT, r2=19, wrap=True, italic=True, valign="top")

# ---- charts row: IRR | utløpsprofil | kjøpesum og egenkapital
sd.row_dimensions[20].height = 10
sec("B", "G", 21, "IRR per holdeperiode")
sec("H", "M", 21, "Utløpsprofil · % av leie")
sec("N", "S", 21, "Kjøpesum og egenkapital (NOK)")
cats = Reference(rt, min_col=2, min_row=51, max_row=53)
bar(sd, cats, [(Reference(rt, min_col=3, min_row=51, max_row=53), "A · re-leie", DEEP),
               (Reference(rt, min_col=4, min_row=51, max_row=53), "B · seksjonssalg", BRIGHT)], '0.0%', ("B", 22, "G", 35))
bar(sd, Reference(rr, min_col=2, min_row=E0, max_row=E1),
    [(Reference(rr, min_col=4, min_row=E0, max_row=E1), "Uten opsjon", DEEP),
     (Reference(rr, min_col=7, min_row=E0, max_row=E1), "Inkl. opsjoner", BRIGHT)], '0%', ("H", 22, "M", 35), gap=40, labels=False)
MN = '#,##0.0,,"m";(#,##0.0,,"m");"–"'
for c in "NOPQRS": sd[f"{c}22"].fill = fill(DEEP)
txt("N", 22, "R", "Post", size=8, bold=True, color=WHITE)
txt("S", 22, "S", "NOK", size=8, bold=True, color=WHITE, align="right")
kj = [("Eiendomsverdi (GAV)", "=Kjopspris", 0), ("− Skatteulempe (latent skatt)", "=-SkatteulempeInn", 0),
      ("± Øvrige tillegg / (fradrag)", "=JustKjopesum", 0), ("Kjøpesum (aksjeoppgjør)", "=Kjopesum", 1),
      ("+ DD / SPA", "=DDSPA", 0), ("+ Megler / working capital", "=Megler+WC", 0)]
kj += ([("+ Etableringskost lån", "=EtabKost*LaanT0", 0)] if REFP else [("+ Kjøpshonorar REQ (mandat)", "=KjopsHon", 0)])
kj += [("Total investering", "=TotInv", 1)]
kj += ([("− Banklån og selgerkreditt", "=-(LaanT0+SKBelop)", 0)] if REFP else [])
kj += [("Egenkapital", "=Egenkapital", 2)]
for n, (lab, f, kind) in enumerate(kj):
    rr_ = 23 + n
    txt("N", rr_, "R", lab, size=8.5, bold=kind > 0)
    txt("S", rr_, "S", f, size=9, fmt=MN, align="right", bold=kind > 0, color=NAVY if kind else TXT)
    for c in "NOPQRS":
        sd[f"{c}{rr_}"].border = Border(bottom=THIN)
        if kind == 1: sd[f"{c}{rr_}"].fill = fill(PALE); sd[f"{c}{rr_}"].border = Border(top=NAVYLINE, bottom=NAVYLINE)
        if kind == 2: sd[f"{c}{rr_}"].fill = fill(PALEGRN); sd[f"{c}{rr_}"].border = Border(top=NAVYLINE, bottom=NAVYLINE)
rr_ = 23 + len(kj)
txt("N", rr_, "S", '="Netto tillegg/(fradrag) vs. eiendomsverdi: "&FIXED((TotInv-Kjopspris)/1000000,1)&" MNOK ("&FIXED((TotInv/Kjopspris-1)*100,1)&" %)"',
    size=7.5, color=GREYTXT, italic=True, r2=rr_ + 1, wrap=True, valign="top")
sd.row_dimensions[36].height = 10

# ---- return bridge (waterfall) + table
from openpyxl.chart import BarChart as _BC, Series as _S
sec("B", "J", 37, f"Avkastningsbro · base case ({LANG} år, scenario A)")
sec("L", "S", 37, "Fra egenkapital til utbetalt (NOK)")
b0, b1 = R["rt_bridge"]
wc = _BC(); wc.type = "col"; wc.grouping = "stacked"; wc.overlap = 100; wc.gapWidth = 45; wc.style = 2
for col, title, color, nf in [(6, "Base", None, None), (9, "Total", NAVY, '#,##0.0,,"m";;;'),
                              (7, "Økning", BRIGHT, '+#,##0.0,,"m";;;'), (8, "Reduksjon", "D64545", '-#,##0.0,,"m";;;')]:
    s_ = _S(Reference(rt, min_col=col, min_row=b0, max_row=b1), title=title)
    if color is None:
        s_.graphicalProperties.noFill = True; s_.graphicalProperties.line.noFill = True
    else:
        s_.graphicalProperties.solidFill = color; s_.graphicalProperties.line.noFill = True
        s_.dLbls = DataLabelList(); s_.dLbls.showVal = True
        s_.dLbls.showSerName = s_.dLbls.showCatName = s_.dLbls.showLegendKey = False
        s_.dLbls.numFmt = nf; s_.dLbls.position = "inEnd"
        s_.dLbls.txPr = _txt(750, "FFFFFF" if color != BRIGHT else NAVY, True)
    wc.series.append(s_)
wc.set_categories(Reference(rt, min_col=10, min_row=b0, max_row=b1))
wc.legend = None; wc.title = None
wc.y_axis.numFmt = '#,##0,,"m"'; wc.y_axis.scaling.min = 0
wc.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="E6E6E6"))
wc.y_axis.spPr = GraphicalProperties(ln=LineProperties(noFill=True))
wc.x_axis.spPr = GraphicalProperties(ln=LineProperties(solidFill="BFBFBF"))
wc.x_axis.delete = False; wc.y_axis.delete = False
wc.x_axis.txPr = _txt(750); wc.y_axis.txPr = _txt(750, "7F7F7F")
wc.x_axis.majorTickMark = "none"; wc.y_axis.majorTickMark = "none"
anchor(wc, "B", 38, "J", 51)
sd.add_chart(wc)
for c in "LMNOPQRS": sd[f"{c}38"].fill = fill(DEEP)
txt("L", 38, "P", "Komponent", size=8, bold=True, color=WHITE)
txt("Q", 38, "R", "NOK", size=8, bold=True, color=WHITE, align="right")
txt("S", 38, "S", "x EK", size=8, bold=True, color=WHITE, align="right")
for n in range(b1 - b0 + 1):
    rr_ = 39 + n; src = b0 + n; tot = n in (0, b1 - b0)
    sd.row_dimensions[rr_].height = 17
    txt("L", rr_, "P", f"=Returns!B{src}", size=8.5, bold=tot)
    txt("Q", rr_, "R", f"=Returns!C{src}", size=9, fmt=MN, align="right", bold=tot, color=NAVY if tot else TXT)
    txt("S", rr_, "S", f"=Returns!C{src}/Egenkapital", size=9, fmt='0.00"x";(0.00"x");"–"', align="right", bold=tot)
    for c in "LMNOPQRS":
        sd[f"{c}{rr_}"].border = Border(bottom=THIN)
        if tot: sd[f"{c}{rr_}"].fill = fill(PALEGRN); sd[f"{c}{rr_}"].border = Border(top=NAVYLINE, bottom=NAVYLINE)
txt("L", 48, "S", "Avdrag er nøytrale (løpende nedbetaling gir tilsvarende lavere restgjeld ved exit). Gevinst = total utbetalt − egenkapital = MoE − 1 × EK.",
    size=7.5, color=GREYTXT, italic=True, r2=50, wrap=True, valign="top")
sd.row_dimensions[52].height = 10

# ---- sensitivity (left) and tenant mix (right)
sec("B", "J", 53, f"Sensitivitet · IRR {LANG} år (A) · KPI × exit yield")
sec("L", "S", 53, "Leietakermiks")
g1, _ = SG[1]
for c in "BCDEFGH": sd[f"{c}54"].fill = fill(DEEP)
txt("B", 54, "C", "KPI \\ exit yield", size=8, bold=True, color=WHITE, align="center")
for j, col in enumerate("DEFGH"):
    txt(col, 54, col, f"=Sensitivities!{get_column_letter(3+j)}16", size=8, bold=True, color=WHITE, align="center", fmt='0.00%')
for i in range(5):
    rr_ = 55 + i
    sd.row_dimensions[rr_].height = 17
    c = txt("B", rr_, "C", f"=Sensitivities!B{17+i}", size=8.5, bold=True, color=WHITE, align="center", fmt='0.0%')
    sd[f"B{rr_}"].fill = fill(DEEP); sd[f"C{rr_}"].fill = fill(DEEP)
    for j, col in enumerate("DEFGH"):
        cell = txt(col, rr_, col, f"={g1[(i, j)].replace('C', 'Sensitivities!C', 0) if False else 'Sensitivities!' + g1[(i, j)]}",
                   size=9, align="center", fmt='0.0%', bold=(i == 2 and j == 2), color=NAVY if (i == 2 and j == 2) else TXT)
        cell.border = Border(bottom=THIN)
sd["F57"].border = Border(left=Side(style="medium", color=NAVY), right=Side(style="medium", color=NAVY),
                          top=Side(style="medium", color=NAVY), bottom=Side(style="medium", color=NAVY))
sd.conditional_formatting.add("D55:H59", ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="7FCBFF"))
txt("I", 55, "J", "Base case uthevet. Full omregning per celle – se Sensitivities for pris-, boligpris/capex- og MoE-grid.",
    size=7.5, color=GREYTXT, r2=59, wrap=True, italic=True, valign="top")

for c in "LMNOPQRS": sd[f"{c}54"].fill = fill(DEEP)
for c1, c2, v, al in [("L", "N", "Leietaker", "left"), ("O", "O", "Kategori", "center"), ("P", "P", "Leie 2026", "center"),
                      ("Q", "Q", "Andel", "center"), ("R", "S", "Kontraktsutløp", "center")]:
    txt(c1, 54, c2, v, size=8, bold=True, color=WHITE, align=al)
mix = [("Oslo kommune", "='Rent roll'!D11", "='Rent roll'!F11", "='Rent roll'!I11"),
       ("Rema 1000 (butikk + garasje)", "='Rent roll'!D7", "='Rent roll'!F7+'Rent roll'!F8", "='Rent roll'!I7"),
       ("Telenor Norge AS", "='Rent roll'!D9", "='Rent roll'!F9", "='Rent roll'!I9"),
       ("Telia Norge AS", "='Rent roll'!D10", "='Rent roll'!F10", "='Rent roll'!I10")]
for n, (lab, cat, rent, exp) in enumerate(mix):
    rr_ = 55 + n
    txt("L", rr_, "N", lab, size=9, bold=True)
    txt("O", rr_, "O", cat, size=8.5, align="center")
    txt("P", rr_, "P", rent, size=9, fmt='#,##0.00,,"m"', align="center")
    txt("Q", rr_, "Q", f"=P{rr_}/Leie2026", size=9, fmt='0.0%', align="center", bold=True, color=NAVY)
    txt("R", rr_, "S", exp, size=9, fmt='mmm yyyy', align="center")
    for c in "LMNOPQRS": sd[f"{c}{rr_}"].border = Border(bottom=THIN)
txt("L", 59, "N", "Sum", size=9, bold=True)
txt("P", 59, "P", "=Leie2026", size=9, fmt='#,##0.00,,"m"', align="center", bold=True)
txt("Q", 59, "Q", "=SUM(Q55:Q58)", size=9, fmt='0.0%', align="center", bold=True)
txt("R", 59, "S", '="WAULT "&FIXED(WAULT,1)&" år"', size=8.5, align="center")
for c in "LMNOPQRS":
    sd[f"{c}59"].fill = fill(PALE); sd[f"{c}59"].border = Border(top=NAVYLINE, bottom=NAVYLINE)
sd.row_dimensions[60].height = 10

# ---- assumptions (left) and value support (right)
sec("B", "J", 61, "Forutsetninger · base case")
sec("L", "S", 61, "Verdistøtte og residualrisiko")
bul = [
    '="Eiendomsverdi "&FIXED(Kjopspris/1000000,0)&" MNOK 1.1."&YEAR(Start)&" · total investering "&FIXED(TotInv/1000000,1)&" MNOK"&IF(KjopsHon>0," inkl. kjøpshonorar "&FIXED(KjopsHonPct*100,1)&" %","")&" · "&IF(LTV=0,"ubelånt","LTV "&FIXED(LTV*100,0)&" %")&" · løpende skatt "&FIXED(SkattLop*100,0)&" %."',
    '="KPI "&FIXED(KPI*100,1)&" % p.a. (Rema 80 %, øvrige 100 %) · 2026→27: "&FIXED(KPIstart*100,1)&" %."',
    ('="Eierkost "&FIXED(Eierkost*100,1)&" % av leie ("&FIXED(EierkostKvm,0)&" kr/kvm år 1) + eiendomsskatt · forvaltning ' + ('(TF, FF; ingen AM i RREF) ' if REFP else '(AM, TF, FF) ') + '"&FIXED(-(CF!E' + str(A['am']) + '+CF!E' + str(A['tf']) + '+CF!E' + str(A['ff']) + ')/1000,0)&" TNOK år 1."'),
    ('="Strukturell ledighet "&FIXED(StruktLed*100,1)&" % fra år "&StruktFra&" · ingen capex-lekkasje i RREF (capex modelleres eksplisitt)."' if REFP else
     '="Strukturell ledighet "&FIXED(StruktLed*100,1)&" % fra år "&StruktFra&" · capex-lekkasje "&FIXED(CapexLek*100,1)&" % av leie."'),
    '="Exit yield "&FIXED(ExitYield*100,2)&" % (= inngangsyield"&IF(ExitDelta=0,")"," + "&FIXED(ExitDelta*10000,0)&" bp)")&" på annualisert NOI år N+1 · divestment "&FIXED(Divest*100,2)&" %."',
    '="Oslo kommune utløper "&YEAR(OKUtlop)&" uten gjenværende opsjon · A: re-leie på markedsleie etter "&\'Rent roll\'!P11&" mnd ledighet."',
    '="B: seksjonssalg "&YEAR(SalgDato)&" à "&FIXED(EgenKvm,0)&" kr/kvm, oppussing "&FIXED(CapexKvm,0)&" kr/kvm, gevinstskatt "&FIXED(SkattSeksjon*100,0)&" %."',
    '="Skatteulempe "&FIXED(SUsats*100,0)&" % × (GAV − saldo − "&FIXED(TomtPct*100,0)&" % tomt) ved kjøp og exit · saldo rullert fra 31.12.2024."',
    ('="Bank "&FIXED(LTV*100,0)&" % LTV · margin "&FIXED(Margin*100,2)&" % + NIBOR ("&FIXED(Financing!E9*100,2)&" % år 1) · avdrag "&FIXED(Avdrag*100,1)&" % (avdragsfri "&IOmnd&" mnd) · selgerkreditt "&FIXED(SKBelop/1000000,1)&" MNOK."' if REFP else
    '="Gule celler i Overview er foreløpige (transaksjonskost, forvaltning, boligpris, capex) – se Modellkontroller: "&ModellOK&"."'),
]
for n, f in enumerate(bul):
    rr_ = 62 + n
    sd.row_dimensions[rr_].height = 16
    c = txt("B", rr_, "J", f, size=8.5)
    txt("B", rr_, "B", None) if False else None
    for col in "BCDEFGHIJ": sd[f"{col}{rr_}"].border = Border(bottom=THIN)
vs = [("Kjøpspris", "=Kjopspris", '#,##0.0,,"m"', False),
      ("− Verdi næringsdel (Rema + telekom @ yield næring)", "=-VerdiNaering*-1", '#,##0.0,,"m"', False),
      ("Implisitt verdi boligdel", "=ImplBolig", '#,##0.0,,"m"', True),
      ("Implisitt pris per kvm bolig", "=ImplKvm", '#,##0', False),
      ("Markedspris/kvm (9 komparable salg i bygget)", "=MarkedKvm", '#,##0', False),
      ("Rabatt vs. markedspris", "=ImplRabatt", '0.0%', True),
      ("Netto residualverdi i dag (etter capex og salgskost)", "=Residual0", '#,##0.0,,"m"', False),
      ("Residual i % av implisitt boligverdi", "=Dekning", '0%', True),
      ("Kjøpspris vs. bottom-up verdi", "=Kjopspris/VerdiBottomUp-1", '0.0%', False)]
for n, (lab, f, fmt, b) in enumerate(vs):
    rr_ = 62 + n
    txt("L", rr_, "Q", lab, size=8.5, bold=b)
    txt("R", rr_, "S", f, size=9, fmt=fmt, align="right", bold=b, color=NAVY if b else TXT)
    for col in "LMNOPQRS": sd[f"{col}{rr_}"].border = Border(bottom=THIN)
sd["R63"].value = "=VerdiNaering"   # VerdiNaering is stored negative on Boligverdi
sd.row_dimensions[71].height = 10

# ---- reconciliation vs pipeline deck (EIF) / return bridge (REF)
if REFP:
    sec("B", "S", 72, f"Avkastningsnivåer og finansiering · base case ({LANG} år, scenario A)")
    for c in COLS: sd[f"{c}73"].fill = fill(DEEP)
    for c1, c2, v, al in [("B", "E", "Nivå / nøkkeltall", "left"), ("F", "G", "IRR / verdi", "center"), ("H", "I", "MoE / krav", "center"),
                          ("J", "K", "Status", "center"), ("L", "S", "Kommentar", "left")]:
        txt(c1, 73, c2, v, size=8, bold=True, color=WHITE, align=al)
    rec = [("Ubelånt (eiendomsnivå)", "=Returns!E64", "=Returns!E65", None, '0.0%', '0.00"x"', "Før finansiering, skatt uten rentefradrag"),
           ("Belånt brutto (deal)", "=Returns!E28", "=Returns!E29", None, '0.0%', '0.00"x"', "Etter renter, avdrag og 22 % skatt"),
           ("Etter forvaltningshonorar", "=Utbyttefordeling!E9", None, None, '0.0%', '0.00"x"', "1,6 % p.a. av investert EK"),
           ("Netto investor (etter honorar og carry)", "=Utbyttefordeling!F9", "=Utbyttefordeling!G9", None, '0.0%', '0.00"x"', "SHA 10.2: 8 % pref, catch-up 50/50, 80/20"),
           ("Carried interest / honorar totalt", "=Utbyttefordeling!H9", "=Utbyttefordeling!I9", None, '#,##0.0,,"m"', '#,##0.0,,"m"', "Carry til founder · honorar til forvalter"),
           ("Maks LTV total i holdeperioden", "=LTVmaks", "=CovLTV", '=IF(LTVmaks<=CovLTV,"OK","BRUDD")', '0.0%', '0.0%', "Covenant-krav foreløpig"),
           ("Min ICR i holdeperioden", "=ICRmin", "=CovICR", '=IF(ICRmin>=CovICR,"OK","BRUDD")', '0.00"x"', '0.00"x"', "EBITDA / rentekostnad"),
           ("Debt yield år 1", "=DY1", None, None, '0.0%', None, "NOI / gjeld"),
           ("Gjeld ved exit", "=GjeldExit", None, None, '#,##0.0,,"m"', None, "Innfris fra exit-proveny"),
           ("Covenant-brudd (antall)", "=CovBrudd", None, '=IF(CovBrudd=0,"OK","BRUDD")', '0', None, "Se Financing · rad 36–37")]
    for n, (lab, f1, f2, st, fm1, fm2, com) in enumerate(rec):
        rr_ = 74 + n
        sd.row_dimensions[rr_].height = 16
        txt("B", rr_, "E", lab, size=8.5, bold=True)
        txt("F", rr_, "G", f1, size=9, fmt=fm1, align="center", color=NAVY, bold=True)
        txt("H", rr_, "I", f2, size=9, fmt=fm2 or "General", align="center")
        txt("J", rr_, "K", st, size=9, align="center", bold=True)
        txt("L", rr_, "S", com, size=7.5, color=GREYTXT, italic=True)
        for col in COLS: sd[f"{col}{rr_}"].border = Border(bottom=THIN)
    sd.conditional_formatting.add("J74:K83", CellIsRule(operator="equal", formula=['"BRUDD"'], font=Font(color="C00000", bold=True)))
    sd.conditional_formatting.add("J74:K83", CellIsRule(operator="equal", formula=['"OK"'], font=Font(color="1E7B34", bold=True)))
else:
    sec("B", "S", 72, "Avstemming mot pipeline deck (17.09.2026)")
    for c in COLS: sd[f"{c}73"].fill = fill(DEEP)
    for c1, c2, v, al in [("B", "E", "Nøkkeltall", "left"), ("F", "G", "Modell", "center"), ("H", "I", "Deck", "center"),
                          ("J", "K", "Avvik", "center"), ("L", "S", "Kommentar", "left")]:
        txt(c1, 73, c2, v, size=8, bold=True, color=WHITE, align=al)
    rec = [("Kjøpspris", "=Kjopspris", 130_000_000, '#,##0.0,,"m"', "Oppdatert bud-nivå"),
           ("Netto yield på kjøp", "=EntryYield", 0.058, '0.00%', "Deck: leie − eierkost (ekskl. eiendomsskatt). Modell: NOI etter eiendomsskatt."),
           ("WAULT", "=WAULT", 8.9, '0.0" år"', "Modell: fra 01.01.2027, uten opsjoner"),
           ("Brutto leie år 1", f"={grossA}", 7_900_000, '#,##0.00,,"m"', "2026-leie indeksert med faktisk KPI 2027"),
           ("Pris/kvm BTA", "=Kjopspris/ArealTot", 49_789, '#,##0', ""),
           ("Offentlig andel av leie", "=OffAndel", 0.749, '0.0%', ""),
           ("IRR 10 år (A)", "=Returns!E28", 0.071, '0.0%', "Revidert: rullert saldo og faktisk tomteverdi, forvaltningskost, XIRR-datering"),
           ("Equity multiple 10 år (A)", "=Returns!E29", 1.77, '0.00"x"', ""),
           ("Cash-on-cash år 3 (A)", "=Returns!C44", 0.056, '0.0%', ""),
           ("DPI år 10 (A)", "=Returns!E45", 0.588, '0.0%', "")]
    for n, (lab, f, deck, fmt, com) in enumerate(rec):
        rr_ = 74 + n
        sd.row_dimensions[rr_].height = 16
        txt("B", rr_, "E", lab, size=8.5, bold=True)
        txt("F", rr_, "G", f, size=9, fmt=fmt, align="center", color=NAVY, bold=True)
        c = txt("H", rr_, "I", deck, size=9, fmt=fmt, align="center", color="0000FF")
        txt("J", rr_, "K", f"=F{rr_}-H{rr_}", size=9, fmt=fmt.replace(';', ';') if fmt != '#,##0.0,,"m"' else '#,##0.0,,"m";(#,##0.0,,"m");"–"', align="center")
        txt("L", rr_, "S", com, size=7.5, color=GREYTXT, italic=True)
        for col in COLS: sd[f"{col}{rr_}"].border = Border(bottom=THIN)


rr_ = 74 + len(rec) + 1
sd.row_dimensions[rr_].height = 18
txt("B", rr_, "H", '="Modellkontroller: "&ModellOK', size=8.5, bold=True, color=DEEP)
txt("I", rr_, "S", "Kilder: leietakeroversikt 2026, leiekontrakter, saldoskjema 31.12.2024, pipeline deck 17.09.2026, markedsdata Røa (se Boligverdi).",
    size=7.5, color=GREYTXT, italic=True, align="right")
sd.conditional_formatting.add(f"B{rr_}", CellIsRule(operator="equal", formula=['"Modellkontroller: OK"'], font=Font(color="1E7B34", bold=True)))
sd.print_area = f"A1:T{rr_}"
sd.page_setup.fitToHeight = 1

# ---- final cosmetics
for rr_ in range(R['rr_T0'], R['rr_T1'] + 2):
    for col in range(2, 19):
        c = rr.cell(row=rr_, column=col)
        c.alignment = Alignment(horizontal=c.alignment.horizontal, vertical="center", wrap_text=c.alignment.wrap_text, indent=c.alignment.indent)
for ws in wb.worksheets:
    ws.sheet_view.selection[0].activeCell = "A1"; ws.sheet_view.selection[0].sqref = "A1"
wb.active = 0
wb.calculation.fullCalcOnLoad = True
OUT = "./" + ("Aslakveien_18_Underwriting_model_REF_belaant.xlsx" if REFP else "Aslakveien_18_Underwriting_model.xlsx")
wb.save(OUT)
print("saved", OUT)
