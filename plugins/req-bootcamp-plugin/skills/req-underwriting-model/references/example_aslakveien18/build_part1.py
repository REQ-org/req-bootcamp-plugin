# Worked example: Aslakveien 18. Copy scripts/req_style.py here. Run: PROFILE=EIF python3 build_part4.py (Evergreen, ubelånt) or PROFILE=REF python3 build_part4.py (Real Estate Fund, belånt).
import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, Border
from openpyxl.utils import get_column_letter, column_index_from_string as ci
from req_style import *
import os
PROFILE = os.environ.get('PROFILE', 'EIF'); REFP = PROFILE == 'REF'; LANG = 7 if REFP else 10

wb = Workbook()
wb.remove(wb.active)
S = {}  # sheet handles

def mk(title, tab):
    ws = wb.create_sheet(title); S[title] = ws
    if tab: ws.sheet_properties.tabColor = tab
    return ws

for t, tab in [("Sammendrag", NAVY), ("Overview", DEEP), ("Rent roll", DEEP), ("Boligverdi", DEEP),
               ("CF", BRIGHT), ("Tax accounts", BRIGHT), ("Financing", BRIGHT), ("Returns", BRIGHT),
               ] + ([("Utbyttefordeling", BRIGHT)] if REFP else []) + [("Sensitivities", BRIGHT), ("Sens-motor", "BFBFBF")]:
    mk(t, tab)

YC = [get_column_letter(5 + i) for i in range(11)]   # E..O = years 1..11 on time-series sheets
R = {}  # row registry

# =====================================================================================
# OVERVIEW (all inputs)
# =====================================================================================
ov = S["Overview"]
setup(ov, {"A": 2, "B": 50, "C": 16, "D": 9, "E": 13, "F": 78}, zoom=90)
banner(ov, "OVERVIEW · FORUTSETNINGER", "Alle input samlet · blå = input · gul = foreløpig/ubekreftet · svart = beregnet · grønn = lenke", "B", "F")
HDR = ["Parameter", "Verdi", "Enhet", "Status", "Kilde / kommentar"]

def inrow(r, label, value, unit, status, note, nm=None, fmt=None, kind=None):
    put(ov, f"B{r}", label, "label")
    k = kind or ("placeholder" if status == "Foreløpig" else ("calc" if isinstance(value, str) and value.startswith("=") else "input"))
    put(ov, f"C{r}", value, k, fmt=fmt)
    put(ov, f"D{r}", unit, "note", align="center")
    put(ov, f"E{r}", status, "text", size=8, align="center")
    put(ov, f"F{r}", note, "note")
    if nm: name(wb, nm, "Overview", f"C{r}")

r = 5
def sec(title):
    global r
    section(ov, r, "B", "F", title); header_row(ov, r + 1, "B", "F", HDR); r += 2

sec("1 · Prosjekt")
for lab, val in [("Eiendom", "Aslakveien 18, 0753 Oslo (Røa)"), ("Segment", "Eldre-/omsorgsbolig og nærsenter"),
                 ("Selskap (SPV)", "Aslakveien 16-18 AS · org.nr 979 992 890"),
                 ("Matrikkel", "0301/11/610, snr 1–28 + 40–41 (allerede seksjonert)"),
                 ("Fond", "REQ Real Estate Fund II (belånt)" if REFP else "REQ Evergreen Income Fund (ubelånt)")]:
    put(ov, f"B{r}", lab, "label"); put(ov, f"C{r}", val, "text"); ov.merge_cells(f"C{r}:D{r}"); r += 1
inrow(r, "Prosjektstart / verdidato (t0)", dt.date(2027, 1, 1), "dato", "Antakelse", "Kjøp antatt 01.01.2027; kontantstrømmer datert 31.12 hvert år", "Start", F_DATE); r += 2

sec("2 · Kjøp og transaksjonskostnader")
inrow(r, "Kjøpspris (eiendomsverdi / GAV)", 130_000_000, "NOK", "Iht. deck", "Oppdatert bud-nivå, pipeline deck 17.09.2026. Driver egenkapital og AM-kostnad.", "Kjopspris", F_NOK); r += 1
inrow(r, "DD / SPA-kostnader", 0, "NOK", "Foreløpig", "Venter på tall", "DDSPA", F_NOK); r += 1
inrow(r, "Working capital", 0, "NOK", "Foreløpig", "Venter på tall", "WC", F_NOK); r += 1
inrow(r, "Megler / øvrige kjøpskostnader", 0, "NOK", "Foreløpig", "Venter på tall", "Megler", F_NOK); r += 2

sec("3 · Skatt")
inrow(r, "Skatteulempe-sats", 0.10, "%", "REQ-standard", "Skatteulempe = sats × (GAV − (saldo + tomteverdi)) ved kjøp og exit", "SUsats", F_PCT); r += 1
inrow(r, "Tomteverdi i skatteulempe-formel, % av GAV", 0.20, "%", "REQ-standard", "Vurderes per case", "TomtPct", F_PCT); r += 1
inrow(r, "Saldodato (skattemessige verdier)", dt.date(2024, 12, 31), "dato", "Iht. saldoskjema", "Saldoskjema i leietakeroversikt (arbeidsbok)", "SaldoDato", F_DATE); r += 1
inrow(r, "Saldo gr. I (bygg) per saldodato", 6_493_255, "NOK", "Iht. saldoskjema", "", "SaldoByggRaw", F_NOK); r += 1
inrow(r, "Saldo gr. J (tekniske anlegg) per saldodato", 1_430_819, "NOK", "Iht. saldoskjema", "", "SaldoTekRaw", F_NOK); r += 1
inrow(r, "Skattemessig inngangsverdi tomt", 3_957_000, "NOK", "Iht. saldoskjema", "Brukes ved faktisk realisasjon (seksjonssalg)", "TomtSkatt", F_NOK); r += 1
inrow(r, "Avskrivningssats gr. I", 0.02, "%", "Skatteloven", "", "AvskrBygg", F_PCT); r += 1
inrow(r, "Avskrivningssats gr. J", 0.10, "%", "Skatteloven", "", "AvskrTek", F_PCT); r += 1
inrow(r, "Saldo gr. I ved prosjektstart (rullert frem)", "=SaldoByggRaw*(1-AvskrBygg)^(YEAR(Start)-1-YEAR(SaldoDato))", "NOK", "Beregnet", "Saldo 31.12.2024 avskrevet 2025–2026", "SaldoBygg", F_NOK); r += 1
inrow(r, "Saldo gr. J ved prosjektstart (rullert frem)", "=SaldoTekRaw*(1-AvskrTek)^(YEAR(Start)-1-YEAR(SaldoDato))", "NOK", "Beregnet", "", "SaldoTek", F_NOK); r += 1
inrow(r, "Løpende skattesats", 0.22 if REFP else 0.0, "%", "Fondsprofil", "22 % med rentefradrag og fremførbart underskudd" if REFP else "0 %-struktur (ubelånt fond). Aksjonærlån kan gi 0–22 % – avventer beregning.", "SkattLop", F_PCT); r += 1
inrow(r, "Fremførbart underskudd ved kjøp", 0, "NOK", "Foreløpig", "Fra selskapets skattemelding", "UnderskuddInn", F_NOK); r += 1
inrow(r, "Gevinstskatt ved seksjonssalg (direkte realisasjon)", 0.22, "%", "Antakelse", "Ordinær selskapsskatt. Gevinst- og tapskonto (20 % p.a.) kan gi utsettelse – ikke hensyntatt (konservativt).", "SkattSeksjon", F_PCT); r += 2

sec("4 · Leieinntekter og indeksering")
inrow(r, "KPI-justering 2026 → 2027", 0.032, "%", "Oppgitt", "Faktisk indeksering for første år, vektet med leietakers KPI-andel", "KPIstart", F_PCT); r += 1
inrow(r, "KPI p.a. fra 2028", 0.02, "%", "Iht. deck", "Pipeline deck: base 2,0 % (sensitivitet 1,5–2,5 %)", "KPI", F_PCT); r += 1
inrow(r, "Strukturell ledighet, % av leie", 0.02, "%", "Iht. deck", "Systematisk ledighet i underwriting", "StruktLed", F_PCT); r += 1
inrow(r, "Strukturell ledighet fra år", 3, "år", "Iht. deck", "", "StruktFra", F_INT); r += 1
inrow(r, "Ledig areal i dag", 0, "kvm", "Iht. rent roll", "Ingen fysisk ledighet", "LedigKvm", F_NOK); r += 2

sec("5 · Driftskostnader, forvaltning og capex")
inrow(r, "Eierkost, % av leie", 0.05, "%", "REQ-standard", "Eiendomsskatt kommer i tillegg", "Eierkost", F_PCT); r += 1
inrow(r, "Eiendomsskatt 2026", 102_854, "NOK", "Iht. rent roll", "Indekseres med KPI", "EiendSkatt", F_NOK); r += 1
inrow(r, "AM-kostnad, % av kjøpspris p.a.", 0.005, "%", "Foreløpig", "Videreført fra Claude Monets Allé 21 – bekreft", "AMpct", F_PCT2); r += 1
inrow(r, "Teknisk forvaltning p.a. (2026-kr)", 80_000, "NOK", "Foreløpig", "Videreført fra Claude Monets Allé 21 – bekreft", "TFkr", F_NOK); r += 1
inrow(r, "Forretningsførsel, % av leie", 0.005, "%", "Foreløpig", "Videreført fra Claude Monets Allé 21 – bekreft", "FFpct", F_PCT2); r += 1
inrow(r, "Capex-lekkasje, % av leie", 0.05, "%", "Iht. deck", "Avsatt utover ordinær reserve (eierkost). Påvirker kontantstrøm, ikke verdivurdering.", "CapexLek", F_PCT); r += 2

sec("6 · Exit")
inrow(r, "Holdeperiode – kort", 3, "år", "Standard", "", "ExitKort", F_INT); r += 1
inrow(r, "Holdeperiode – mellom", 5, "år", "Standard", "", "ExitMellom", F_INT); r += 1
inrow(r, "Holdeperiode – lang (base case)", LANG, "år", "Fondsprofil", "Real Estate Fund: 3/5/7 år (bekreft)" if REFP else "Base case i pipeline deck", "ExitLang", F_INT); r += 1
R["ov_entry"] = r
inrow(r, "Netto yield ved kjøp (NOI år 1 / kjøpspris)", "=CF!E{NOI_A}/Kjopspris", "%", "Beregnet", "NOI etter eierkost og eiendomsskatt", "EntryYield", F_PCT2); r += 1
inrow(r, "Endring exit yield vs. inngang", 0.0, "%-p", "Iht. deck", "Deck: exit yield = inngangsyield", "ExitDelta", F_PCT2); r += 1
inrow(r, "Exit yield (scenario A)", "=EntryYield+ExitDelta", "%", "Beregnet", "Anvendes på annualisert NOI år N+1", "ExitYield", F_PCT2); r += 1
inrow(r, "Yield næringsdel (Rema + telekom)", 0.065, "%", "Antakelse", "Scenario B (gjenværende næring) og verdistøtte. Iht. REQ arbeidsbok.", "YieldNaering", F_PCT2); r += 1
inrow(r, "Divestment-kost, % av exit-verdi", 0.0125, "%", "REQ-standard", "", "Divest", F_PCT2); r += 2

sec("7 · Scenario B – Oslo kommune fraflytter, seksjonssalg")
inrow(r, "Kontraktsutløp Oslo kommune", "='Rent roll'!I11", "dato", "Lenke", "Ingen gjenværende opsjon", "OKUtlop", F_DATE, kind="link"); r += 1
inrow(r, "Seksjonssalg, måneder etter utløp", 12, "mnd", "Antakelse", "Tid til oppussing og salg av 28 enheter", "SalgMnd", F_INT); r += 1
inrow(r, "Salgsdato seksjoner", "=EDATE(OKUtlop,SalgMnd)", "dato", "Beregnet", "", "SalgDato", F_DATE); r += 1
R["ov_andel"] = r
inrow(r, "Boligdelens andel (allokeringsnøkkel)", "=CF!E{RRNB_A}/CF!E{RRN_A}", "%", "Beregnet", "NOI-andel år 1 – fordeler saldo, tomt, eiendomsskatt og AM ved seksjonssalg", "BoligAndel", F_PCT); r += 1
inrow(r, "Boligprisvekst p.a.", 0.02, "%", "Antakelse", "Satt lik KPI", "BoligVekst", F_PCT); r += 1
inrow(r, "Egen pris/kvm boligdel i salgbar stand (2026-kr)", 92_461, "NOK/kvm", "Foreløpig", "Nøytralt satt lik markedspris (komparable) – settes av REQ", "EgenKvm", F_NOK); r += 1
inrow(r, "Oppussingscapex til salgbar stand (2026-kr)", 10_000, "NOK/kvm", "Foreløpig", "Ingen kartlegging gjort – generelle anslag 5–15 000 kr/kvm. Hent tilbud.", "CapexKvm", F_NOK); r += 1
inrow(r, "Salgbart areal bolig (BRA)", 1_596, "kvm", "Iht. rent roll", "BRA plan 3–4 (798 + 798). Optimo-spesifikasjon: 26 enheter ≈ 1 509 kvm + 2 fellesrom ≈ 125 kvm.", "ArealSalg", F_NOK); r += 1
inrow(r, "Salgskostnad seksjoner", 0.02, "%", "Antakelse", "Megler og markedsføring", "SalgKost", F_PCT); r += 1
inrow(r, "P-plasser solgt med seksjonene", 0, "stk", "Antakelse", "Oppside: 22 plasser (12 garasje + 10 ute), men 20 er knyttet til Rema-kontrakten", "PPlasser", F_INT); r += 1
inrow(r, "Verdi per p-plass (2026-kr)", 400_000, "NOK", "Iht. arbeidsbok", "", "PVerdi", F_NOK); r += 2

sec("8 · Verdistøtte (bottom-up)")
inrow(r, "Yield boligdel – going concern (offentlig leietaker)", 0.0525, "%", "Antakelse", "Kun verdistøtte – påvirker ikke avkastning", "CapBolig", F_PCT2); r += 2

sec("9 · Finansiering")
inrow(r, "Belåningsgrad bank (LTV av kjøpspris)", 0.60 if REFP else 0.0, "%", "Fondsprofil", "REQ Real Estate Fund: 60 %" if REFP else "Fondet er ubelånt – sett LTV > 0 for belånt analyse", "LTV", F_PCT); r += 1
inrow(r, "Margin", 0.02 if REFP else 0.0175, "%", "Fondsprofil", "Over 3M NIBOR (forward-kurve under)", "Margin", F_PCT2); r += 1
inrow(r, "Avdrag p.a., % av opprinnelig lån", 0.02, "%", "Fondsprofil", "Betales løpende etter avdragsfri periode", "Avdrag", F_PCT); r += 1
inrow(r, "Avdragsfri periode", 0, "mnd", "Antakelse", "0 = avdrag fra første periode", "IOmnd", F_INT); r += 1
inrow(r, "Første avdragsdato", "=EDATE(Start,IOmnd)", "dato", "Beregnet", "", "FirstAmort", F_DATE); r += 1
inrow(r, "Etableringskost, % av lån", 0.003 if REFP else 0.005, "%", "Antakelse", "Betales ved opptrekk (t0)", "EtabKost", F_PCT2); r += 1
inrow(r, "Covenant – maks LTV (total)", 0.65, "%", "Foreløpig", "Bekreft mot term sheet", "CovLTV", F_PCT); r += 1
inrow(r, "Covenant – min ICR (EBITDA / renter)", 2.0, "x", "Foreløpig", "Bekreft mot term sheet", "CovICR", F_X); r += 1
put(ov, f"B{r}", "3M NIBOR forward per år", "label", bold=True); put(ov, f"F{r}", "Foreløpig kurve (3M NIBOR snitt 2025: 4,53 %) – erstatt med bankens forward-kurve", "note"); r += 1
nib0 = r
for k, v in enumerate([0.041, 0.039, 0.038, 0.0375, 0.0375, 0.0375, 0.0375, 0.0375, 0.0375, 0.0375, 0.0375]):
    put(ov, f"B{r}", f"=YEAR(Start)+{k}", "calc", F_INT, align="left"); put(ov, f"C{r}", v, "placeholder", F_PCT2); r += 1
line_rows(ov, nib0, r - 1, "B", "F")
wb.defined_names["NiborTab"] = DefinedName("NiborTab", attr_text=f"'Overview'!$B${nib0}:$C${r-1}")
r += 1
sec("9b · Selgerkreditt")
inrow(r, "Selgerkreditt – beløp", 0, "NOK", "Antakelse", "0 = ingen. Reduserer egenkapital t0; innfris ved forfall eller exit.", "SKBelop", F_NOK); r += 1
inrow(r, "Selgerkreditt – rente (fast)", 0.06, "%", "Antakelse", "", "SKRente", F_PCT2); r += 1
inrow(r, "Selgerkreditt – rentebetaling", "PIK", "PIK/Løpende", "Antakelse", "PIK = kapitaliseres og betales ved innfrielse", "SKPIK"); r += 1
inrow(r, "Selgerkreditt – forfall", dt.date(2029, 12, 31), "dato", "Antakelse", "", "SKForfall", F_DATE); r += 2
if REFP:
    sec("9c · Fond – utbyttefordeling (SHA pkt 10.2) og honorar")
    inrow(r, "Preferred return (rentes rente 31.12)", 0.08, "%", "Iht. SHA", "Pkt 10.2 (a)", "Pref", F_PCT); r += 1
    inrow(r, "Catch-up – andel til founder", 0.50, "%", "Iht. SHA", "Pkt 10.2 (c)", "CatchUp", F_PCT); r += 1
    inrow(r, "Carried interest", 0.20, "%", "Iht. SHA", "Pkt 10.2 (d)", "Carry", F_PCT); r += 1
    inrow(r, "Forvaltningshonorar p.a., % av investert EK", 0.016, "%", "Iht. SHA", "Pkt 14.3 – Employed Capital", "MgmtFee", F_PCT2); r += 2
section(ov, r, "B", "F", "10 · Kilder og anvendelse (t0)"); r += 1
header_row(ov, r, "B", "F", ["Post", "NOK", "", "", "Kommentar"]); r += 1
su0 = r
for lab, f, note in [("Kjøpspris (GAV)", "=Kjopspris", ""),
                     ("− Skatteulempe (latent skatt, trekkes fra pris)", "=-SkatteulempeInn", "Se Tax accounts"),
                     ("+ DD / SPA", "=DDSPA", ""), ("+ Working capital", "=WC", ""), ("+ Megler / øvrige", "=Megler", ""),
                     ("+ Etableringskost lån", "=EtabKost*LaanT0", "")]:
    put(ov, f"B{r}", lab, "label"); put(ov, f"C{r}", f, "calc", F_NOK); put(ov, f"F{r}", note, "note"); r += 1
put(ov, f"B{r}", "Sum anvendelse", "label", bold=True); put(ov, f"C{r}", f"=SUM(C{su0}:C{r-1})", "calc", F_NOK, bold=True)
total_row(ov, r, "B", "F"); R["ov_uses"] = r; r += 1
put(ov, f"B{r}", "Banklån", "label"); put(ov, f"C{r}", "=LTV*Kjopspris", "calc", F_NOK); name(wb, "LaanT0", "Overview", f"C{r}"); r += 1
put(ov, f"B{r}", "Selgerkreditt", "label"); put(ov, f"C{r}", "=SKBelop", "calc", F_NOK); r += 1
put(ov, f"B{r}", "Egenkapital", "label"); put(ov, f"C{r}", f"=C{R['ov_uses']}-LaanT0-SKBelop", "calc", F_NOK); name(wb, "Egenkapital", "Overview", f"C{r}"); r += 1
put(ov, f"B{r}", "Sum kilder", "label", bold=True); put(ov, f"C{r}", "=LaanT0+SKBelop+Egenkapital", "calc", F_NOK, bold=True)
total_row(ov, r, "B", "F"); R["ov_src"] = r; r += 2
R["ov_checks"] = r  # checks written in part 2
ov.freeze_panes = "B5"

# =====================================================================================
# RENT ROLL
# =====================================================================================
rr = S["Rent roll"]
setup(rr, {"A": 2, "B": 22, "C": 24, "D": 10, "E": 9, "F": 12, "G": 9, "H": 8, "I": 11, "J": 8, "K": 11,
           "L": 9, "M": 9, "N": 8, "O": 12, "P": 9, "Q": 10, "R": 58}, zoom=85)
banner(rr, "RENT ROLL · LEIEFORHOLD OG UTLØPSPROFIL", "Kontraktsleie 2026-nivå · markedsleie/ledighet/capex brukes ved kontraktsutløp i kontantstrømmen", "B", "R")
section(rr, 5, "B", "R", "Leietakere")
header_row(rr, 6, "B", "R", ["Leietaker", "Leieobjekt", "Kategori", "Areal BTA (kvm)", "Leie 2026 (NOK)", "Leie/kvm",
                              "KPI-andel", "Kontrakts-utløp", "Opsjon (år)", "Utløp inkl. opsjon", "Gj. løpetid (år)",
                              "Gj. løpetid inkl. ops.", "Andel av leie", "Markedsleie v/utløp (2026-kr)", "Ledighet v/utløp (mnd)",
                              "Capex re-letting (kr/kvm)", "Kommentar"], height=40)
tenants = [
    ("Rema 1000", "Dagligvarebutikk, plan 1", "Privat", 998, 1_829_396, 0.8, dt.date(2041, 4, 30), 5, 1_829_396, 0, 0,
     "80 % KPI. 20 reserverte p-plasser inngår. Utløp utenfor modellhorisont."),
    ("Rema 1000", "Garasje", "Privat", 0, 26_312, 0.8, dt.date(2041, 4, 30), 0, 26_312, 0, 0,
     "Areal inngår i butikk. Følger hovedkontrakt."),
    ("Telenor Norge AS", "Basestasjon", "Privat", 11, 30_174, 1.0, dt.date(2031, 1, 31), 10, 30_174, 3, 0,
     "Markedsleie/ledighet foreløpig"),
    ("Telia Norge AS", "Basestasjon", "Privat", 6, 31_536, 1.0, dt.date(2032, 4, 30), 10, 31_536, 3, 0,
     "Markedsleie/ledighet foreløpig"),
    ("Oslo kommune", "28 omsorgsboliger, plan 3–4", "Offentlig", 1_596, 5_718_674, 1.0, dt.date(2033, 5, 31), 0, 5_718_674, 6, 0,
     "Bydel Røa + Vinderen. 20 år + 10 års opsjon benyttet – ingen gjenværende opsjon (vurdert mot avtaletekst; leietakeroversikt oppga feilaktig opsjon til 2043)."),
]
T0, T1 = 7, 11
for i, t in enumerate(tenants):
    r = T0 + i
    nm_, obj, cat, area, rent, kpi, exp, opt, mkt, vac, capex, note = t
    put(rr, f"B{r}", nm_, "label", bold=True); put(rr, f"C{r}", obj, "text"); put(rr, f"D{r}", cat, "input", align="center")
    put(rr, f"E{r}", area, "input", F_NOK); put(rr, f"F{r}", rent, "input", F_NOK)
    put(rr, f"G{r}", f'=IF(E{r}>0,F{r}/E{r},"–")', "calc", F_NOK)
    put(rr, f"H{r}", kpi, "input", '0%'); put(rr, f"I{r}", exp, "input", F_DATE); put(rr, f"J{r}", opt, "input", F_INT)
    put(rr, f"K{r}", f"=EDATE(I{r},12*J{r})", "calc", F_DATE)
    put(rr, f"L{r}", f"=MAX(0,(I{r}-Start)/365.25)", "calc", '0.0')
    put(rr, f"M{r}", f"=MAX(0,(K{r}-Start)/365.25)", "calc", '0.0')
    put(rr, f"N{r}", f"=F{r}/$F${T1+1}", "calc", F_PCT)
    put(rr, f"O{r}", mkt, "placeholder", F_NOK); put(rr, f"P{r}", vac, "placeholder", F_INT); put(rr, f"Q{r}", capex, "placeholder", F_NOK)
    put(rr, f"R{r}", note, "note")
    rr[f"R{r}"].alignment = Alignment(wrap_text=True, vertical="center")
    rr.row_dimensions[r].height = 30 if len(note) > 60 else 18
line_rows(rr, T0, T1, "B", "R")
r = T1 + 1
put(rr, f"B{r}", "Sum / vektet", "label", bold=True)
put(rr, f"E{r}", f"=SUM(E{T0}:E{T1})", "calc", F_NOK); put(rr, f"F{r}", f"=SUM(F{T0}:F{T1})", "calc", F_NOK)
put(rr, f"G{r}", f"=F{r}/E{r}", "calc", F_NOK)
put(rr, f"L{r}", f"=SUMPRODUCT(F{T0}:F{T1},L{T0}:L{T1})/F{r}", "calc", '0.0')
put(rr, f"M{r}", f"=SUMPRODUCT(F{T0}:F{T1},M{T0}:M{T1})/F{r}", "calc", '0.0')
put(rr, f"N{r}", f"=SUM(N{T0}:N{T1})", "calc", F_PCT)
total_row(rr, r, "B", "R")
TT = r
for nm_, col in [("Leie2026", "F"), ("ArealLeid", "E"), ("WAULT", "L"), ("WAULTopt", "M")]:
    name(wb, nm_, "Rent roll", f"{col}{TT}")

# key metrics
r = TT + 2
section(rr, r, "B", "F", "Nøkkeltall leieforhold"); r += 1
km = [("Totalt areal BTA (kvm)", "=ArealLeid+LedigKvm", F_NOK, "ArealTot"),
      ("Ledighet i bygget (areal)", "=LedigKvm/ArealTot", F_PCT, "Ledighet"),
      ("Offentlig andel av leie", f'=SUMIF(D{T0}:D{T1},"Offentlig",F{T0}:F{T1})/Leie2026', F_PCT, "OffAndel"),
      ("Antall leietakere", f"=SUMPRODUCT(1/COUNTIF(B{T0}:B{T1},B{T0}:B{T1}))", F_INT, "AntLeietakere"),
      ("WAULT uten opsjoner (fra t0)", "=WAULT", '0.0" år"', None),
      ("WAULT inkl. opsjoner (fra t0)", "=WAULTopt", '0.0" år"', None)]
for lab, f, fmt, nm_ in km:
    put(rr, f"B{r}", lab, "label"); rr.merge_cells(f"B{r}:D{r}")
    put(rr, f"E{r}", f, "calc", fmt); rr.merge_cells(f"E{r}:F{r}")
    if nm_: name(wb, nm_, "Rent roll", f"E{r}")
    r += 1
line_rows(rr, r - len(km), r - 1, "B", "F")

# expiry profile
r += 1
section(rr, r, "B", "H", "Utløpsprofil · kontraktsleie som utløper per år (2026-nivå)"); R["rr_exp_sec"] = r; r += 1
header_row(rr, r, "B", "H", ["År", "Uten opsjon (NOK)", "% av leie", "Akk. %", "Inkl. opsjoner (NOK)", "% av leie", "Akk. %"], height=30)
r += 1
E0 = r
for k in range(20):
    rr[f"B{r}"] = f"=YEAR(Start)+{k}"; rr[f"B{r}"].font = font(9, bold=True); rr[f"B{r}"].number_format = F_INT
    rr[f"B{r}"].alignment = Alignment(horizontal="left", indent=1)
    put(rr, f"C{r}", f"=SUMPRODUCT((YEAR($I${T0}:$I${T1})=B{r})*$F${T0}:$F${T1})", "calc", F_NOK)
    put(rr, f"D{r}", f"=C{r}/Leie2026", "calc", F_PCT)
    put(rr, f"E{r}", f"=SUM($D${E0}:D{r})", "calc", F_PCT)
    put(rr, f"F{r}", f"=SUMPRODUCT((YEAR($K${T0}:$K${T1})=B{r})*$F${T0}:$F${T1})", "calc", F_NOK)
    put(rr, f"G{r}", f"=F{r}/Leie2026", "calc", F_PCT)
    put(rr, f"H{r}", f"=SUM($G${E0}:G{r})", "calc", F_PCT)
    r += 1
E1 = r - 1
line_rows(rr, E0, E1, "B", "H")
put(rr, f"B{r}", "Senere", "label", bold=True)
put(rr, f"C{r}", f"=SUMPRODUCT((YEAR($I${T0}:$I${T1})>B{E1})*$F${T0}:$F${T1})", "calc", F_NOK)
put(rr, f"D{r}", f"=C{r}/Leie2026", "calc", F_PCT)
put(rr, f"F{r}", f"=SUMPRODUCT((YEAR($K${T0}:$K${T1})>B{E1})*$F${T0}:$F${T1})", "calc", F_NOK)
put(rr, f"G{r}", f"=F{r}/Leie2026", "calc", F_PCT); r += 1
put(rr, f"B{r}", "Sum", "label", bold=True)
put(rr, f"C{r}", f"=SUM(C{E0}:C{r-1})", "calc", F_NOK); put(rr, f"D{r}", f"=SUM(D{E0}:D{r-1})", "calc", F_PCT)
put(rr, f"F{r}", f"=SUM(F{E0}:F{r-1})", "calc", F_NOK); put(rr, f"G{r}", f"=SUM(G{E0}:G{r-1})", "calc", F_PCT)
total_row(rr, r, "B", "H")
R.update(rr_T0=T0, rr_T1=T1, rr_TT=TT, rr_E0=E0, rr_E1=E1, rr_Esum=r)
rr.freeze_panes = "C7"

# =====================================================================================
# TIME-SERIES SHEET HELPER
# =====================================================================================
def ts_setup(ws, title, subtitle):
    setup(ws, {"A": 2, "B": 46, "C": 10, "D": 2, **{c: 12.5 for c in YC}, "P": 2}, zoom=85)
    banner(ws, title, subtitle, "B", "O")
    years_header(ws, "E", 11, extra_last=True)
    ws["C5"].value = None
    ws.freeze_panes = "E7"

def tsrow(ws, r, label, fn, fmt=F_NOK, kind="calc", bold=False, unit=""):
    put(ws, f"B{r}", label, "label", bold=bold)
    if unit: put(ws, f"C{r}", unit, "note", align="center")
    for i, c in enumerate(YC):
        v = fn(c, i) if callable(fn) else fn
        cell = put(ws, f"{c}{r}", v, kind, fmt, bold=bold)
        if i == 10:
            cell.font = Font(name=cell.font.name, size=cell.font.size, bold=bold, color="8C8C8C", italic=True)

def subhead(ws, r, text):
    c = ws[f"B{r}"]; c.value = text; c.font = Font(name=HEAD, size=8.5, bold=True, color=DEEP)

# =====================================================================================
# TAX ACCOUNTS
# =====================================================================================
tx = S["Tax accounts"]
ts_setup(tx, "TAX ACCOUNTS · SKATTEMESSIGE SALDOER OG SKATTEULEMPE", "Saldo rullert frem fra 31.12.2024 · kolonne 11 kun for fremoverskuende verdier")
section(tx, 8, "B", "O", "Skattemessige saldoer")
prev = lambda c, i, row: (f"={YC[i-1]}{row+2}" if i else None)
tsrow(tx, 9, "Saldo gr. I (bygg) – IB", lambda c, i: "=SaldoBygg" if i == 0 else f"={YC[i-1]}11")
tsrow(tx, 10, "Avskrivning gr. I", lambda c, i: f"={c}9*AvskrBygg")
tsrow(tx, 11, "Saldo gr. I – UB", lambda c, i: f"={c}9-{c}10", bold=True)
tsrow(tx, 12, "Saldo gr. J (tekniske anlegg) – IB", lambda c, i: "=SaldoTek" if i == 0 else f"={YC[i-1]}14")
tsrow(tx, 13, "Avskrivning gr. J", lambda c, i: f"={c}12*AvskrTek")
tsrow(tx, 14, "Saldo gr. J – UB", lambda c, i: f"={c}12-{c}13", bold=True)
tsrow(tx, 15, "Sum skattemessige avskrivninger", lambda c, i: f"={c}10+{c}13", bold=True)
tsrow(tx, 16, "Sum skattemessig saldo – UB", lambda c, i: f"={c}11+{c}14", bold=True)
total_row(tx, 16, "B", "O")
line_rows(tx, 9, 15, "B", "O")
R.update(tx_avskr=15, tx_saldo=16)
section(tx, 18, "B", "O", "Løpende skatt (til informasjon – beregnes i CF)")
R["tx_run"] = 19   # rows 19-22 filled in part 2 (need CF rows)
section(tx, 24, "B", "F", "Skatteulempe ved kjøp")
for j, (lab, f) in enumerate([("Kjøpspris (GAV)", "=Kjopspris"), ("− Skattemessig saldo ved kjøp", "=-(SaldoBygg+SaldoTek)"),
                              ("− Tomteverdi (sats × GAV)", "=-TomtPct*Kjopspris"), ("Grunnlag", "=SUM(E25:E27)"),
                              ("Skatteulempe (sats × grunnlag)", "=SUsats*MAX(0,E28)")]):
    rr_ = 25 + j
    put(tx, f"B{rr_}", lab, "label", bold=(j >= 3)); put(tx, f"E{rr_}", f, "calc", F_NOK, bold=(j >= 3))
line_rows(tx, 25, 28, "B", "E"); total_row(tx, 29, "B", "E")
name(wb, "SkatteulempeInn", "Tax accounts", "E29")
put(tx, "F29", "Trekkes fra kjøpspris i egenkapitalgrunnlaget", "note")

# =====================================================================================
# FINANCING
# =====================================================================================
fi = S["Financing"]
ts_setup(fi, "FINANCING · BANKLÅN, SELGERKREDITT OG COVENANTS", "Rente på snittsaldo × (forward NIBOR + margin) · avdrag fra første avdragsdato · all gjeld innfris ved exit")
YSf = lambda c: f"DATE({c}$6,1,1)"; YEf = lambda c: f"DATE({c}$6,12,31)"; Df = lambda c: f"({YEf(c)}-{YSf(c)}+1)"
section(fi, 8, "B", "O", "Banklån")
tsrow(fi, 9, "3M NIBOR (forward)", lambda c, i: f"=VLOOKUP({c}$6,NiborTab,2,FALSE)", F_PCT2)
tsrow(fi, 10, "Rente bank (NIBOR + margin)", lambda c, i: f"={c}9+Margin", F_PCT2)
tsrow(fi, 11, "Banklån – IB", lambda c, i: "=LaanT0" if i == 0 else f"={YC[i-1]}13")
tsrow(fi, 12, "Avdrag", lambda c, i: f"=-MIN({c}11,Avdrag*LaanT0*MAX(0,{YEf(c)}-MAX({YSf(c)},FirstAmort)+1)/{Df(c)})")
tsrow(fi, 13, "Banklån – UB", lambda c, i: f"={c}11+{c}12", bold=True)
tsrow(fi, 14, "Renter bank", lambda c, i: f"=-AVERAGE({c}11,{c}13)*{c}10")
section(fi, 16, "B", "O", "Selgerkreditt")
fbf = lambda c: f"MAX(0,MIN({YEf(c)},SKForfall)-{YSf(c)}+1)/{Df(c)}"
tsrow(fi, 17, "Selgerkreditt – IB", lambda c, i: "=SKBelop" if i == 0 else f"={YC[i-1]}20")
tsrow(fi, 18, "Kapitalisert rente (PIK)", lambda c, i: f'=IF(SKPIK="PIK",{c}17*SKRente*{fbf(c)},0)')
tsrow(fi, 19, "Innfrielse ved forfall", lambda c, i: f"=IF(YEAR(SKForfall)={c}$6,-({c}17+{c}18),0)")
tsrow(fi, 20, "Selgerkreditt – UB", lambda c, i: f"={c}17+{c}18+{c}19", bold=True)
tsrow(fi, 21, "Kontantrente selgerkreditt", lambda c, i: f'=IF(SKPIK="PIK",0,-{c}17*SKRente*{fbf(c)})')
section(fi, 23, "B", "O", "Sum finansiering")
tsrow(fi, 24, "Renter i kontantstrøm", lambda c, i: f"={c}14+{c}21")
tsrow(fi, 25, "Avdrag og innfrielser", lambda c, i: f"={c}12+{c}19")
tsrow(fi, 26, "Rentekostnad skattemessig (inkl. PIK)", lambda c, i: f"={c}14+{c}21-{c}18")
tsrow(fi, 27, "Total gjeld – UB", lambda c, i: f"={c}13+{c}20", bold=True)
for rr_ in (13, 20, 27): total_row(fi, rr_, "B", "O")
line_rows(fi, 9, 12, "B", "O"); line_rows(fi, 14, 14, "B", "O"); line_rows(fi, 17, 19, "B", "O"); line_rows(fi, 21, 21, "B", "O"); line_rows(fi, 24, 26, "B", "O")
section(fi, 29, "B", "O", "Nøkkeltall og covenants (scenario A)")
R.update(fi_int=24, fi_amort=25, fi_inttax=26, fi_ub=27, fi_bint=14, fi_bamort=12, fi_bib=11, fi_bub=13, fi_skub=20)

import pickle
pickle.dump(R, open("./R.pkl", "wb"))
